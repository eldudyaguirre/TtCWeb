from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.forms import AuthenticationForm, PasswordChangeForm
from django.conf import settings
from django.db import IntegrityError, connection, connections, transaction
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from io import BytesIO
from datetime import datetime
from django.views.decorators.http import require_POST

from .models import (
    Archivo,
    Establecimiento,
    ParametrosCliente,
    PuntoEmision,
    SecuencialDocumento,
    UsuarioCliente,
)
from .services.acceso import obtener_resumen_cliente, validar_acceso_cliente
from .services.archivos import eliminar_archivo, guardar_archivo
from .services.legales import obtener_aceptaciones_pendientes, versiones_legales_vigentes


TEMPLATE_PREVIEWS = {
    'page-404-1': 'page-404-1.html',
    'page-about-1': 'page-about-1.html',
    'page-blog-1': 'page-blog-1.html',
    'page-blog-2': 'page-blog-2.html',
    'page-contact-1': 'page-contact-1.html',
    'page-contact-2': 'page-contact-2.html',
    'page-our-people-1': 'page-our-people-1.html',
    'page-partners-1': 'page-partners-1.html',
    'page-pricing-table-1': 'page-pricing-table-1.html',
    'page-pricing-table-2': 'page-pricing-table-2.html',
    'page-projects-1': 'page-projects-1.html',
    'page-search-1': 'page-search-1.html',
    'page-services-1': 'page-services-1.html',
    'page-single-post-1': 'page-single-post-1.html',
    'page-single-project-1': 'page-single-project-1.html',
    'page-single-service-1': 'page-single-service-1.html',
    'page-testimonials-1': 'page-testimonials-1.html',
    'quienes': 'quienes.html',
}


def home(request):
    return render(request, 'index.html')


def do_signin(request):
    if request.user.is_authenticated:
        return redirect('portal')

    form = AuthenticationForm(request, data=request.POST or None)

    if request.method == 'POST':
        if form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']
            user = authenticate(request, username=username, password=password)

            if user is not None:
                if user.is_staff or user.is_superuser:
                    login(request, user)
                    return redirect('portal')

                asignaciones = UsuarioCliente.objects.filter(
                    usuario=user,
                    activo=True,
                ).select_related('cliente')

                if not asignaciones.exists():
                    messages.error(
                        request,
                        'Tu usuario no tiene un cliente activo asignado al portal.',
                    )
                else:
                    acceso_permitido = False
                    ultimo_motivo = None

                    for asignacion in asignaciones:
                        permitido, _, detalle = validar_acceso_cliente(
                            asignacion.cliente
                        )

                        if permitido:
                            request.session['cliente_ruccedcli'] = (
                                asignacion.cliente.ruccedcli
                            )
                            login(request, user)
                            return redirect('portal')

                        ultimo_motivo = detalle

                    messages.error(
                        request,
                        ultimo_motivo or 'El acceso al portal está restringido.',
                    )
        else:
            messages.error(request, 'Usuario o contraseña inválidos.')

    return render(request, 'sign-in.html', {'signin_form': form})


@login_required
def mi_empresa(request):
    if request.user.is_staff or request.user.is_superuser:
        return redirect('portal')

    asignacion = (
        UsuarioCliente.objects
        .filter(usuario=request.user, activo=True)
        .select_related('cliente')
        .first()
    )

    if asignacion is None:
        logout(request)
        messages.error(
            request,
            'No tienes una empresa activa asignada a tu usuario.',
        )
        return redirect('signin')

    permitido, _, detalle = validar_acceso_cliente(asignacion.cliente)

    if not permitido:
        logout(request)
        messages.error(request, detalle)
        return redirect('signin')

    cliente = asignacion.cliente

    if request.method == 'POST':
        if asignacion.rol == UsuarioCliente.Rol.CONSULTA:
            messages.error(request, 'Tu usuario tiene permisos de consulta y no puede modificar los datos.')
        else:
            campos_editables = [
                'dirclient', 'ocuclient', 'corelectr',
                'teldomcli', 'teloficli', 'telcelcli',
                'estcivcli', 'fecnacimi',
                'clavesri', 'cediess', 'claveiess',
                'mrlcon', 'mrlsal', 'clavesuper',
                'iessdomestica', 'datiess',
            ]

            for campo in campos_editables:
                setattr(cliente, campo, request.POST.get(campo, '').strip())

            cliente.save(update_fields=campos_editables)
            messages.success(request, 'Los datos de la empresa fueron actualizados correctamente.')

        return redirect('mi_empresa')

    return render(
        request,
        'mi-empresa.html',
        {
            'cliente': cliente,
            'rol': asignacion.rol,
            'puede_editar': asignacion.rol != UsuarioCliente.Rol.CONSULTA,
        },
    )


@login_required
def parametros(request):
    if request.user.is_staff or request.user.is_superuser:
        return redirect('portal')

    asignacion = (
        UsuarioCliente.objects
        .filter(usuario=request.user, activo=True)
        .select_related('cliente')
        .first()
    )

    if asignacion is None:
        logout(request)
        messages.error(request, 'No tienes una empresa activa asignada a tu usuario.')
        return redirect('signin')

    permitido, _, detalle = validar_acceso_cliente(asignacion.cliente)
    if not permitido:
        logout(request)
        messages.error(request, detalle)
        return redirect('signin')

    cliente = asignacion.cliente
    puede_editar = asignacion.rol != UsuarioCliente.Rol.CONSULTA
    parametros_cliente, _ = ParametrosCliente.objects.get_or_create(cliente=cliente)

    if request.method == 'POST':
        if not puede_editar:
            messages.error(request, 'Tu usuario tiene permisos de consulta y no puede modificar los parámetros.')
            return redirect('parametros')

        accion = request.POST.get('accion', '').strip()

        try:
            with transaction.atomic():
                if accion == 'guardar_certificado':
                    ambiente = request.POST.get('ambiente', ParametrosCliente.Ambiente.PRUEBAS)
                    if ambiente not in dict(ParametrosCliente.Ambiente.choices):
                        raise ValueError('El ambiente seleccionado no es válido.')

                    parametros_cliente.ambiente = ambiente
                    clave_p12 = request.POST.get('clave_p12', '').strip()
                    if clave_p12:
                        parametros_cliente.clave_p12 = clave_p12

                    archivo = request.FILES.get('certificado_p12')
                    if archivo:
                        if not archivo.name.lower().endswith('.p12'):
                            raise ValueError('El certificado debe ser un archivo con extensión .p12.')
                        if archivo.size > 10 * 1024 * 1024:
                            raise ValueError('El archivo P12 no puede superar los 10 MB.')

                        archivo_anterior = parametros_cliente.certificado_archivo
                        nuevo_archivo = guardar_archivo(
                            cliente=cliente,
                            uploaded_file=archivo,
                            tipo=Archivo.Tipo.CERTIFICADO_P12,
                            subcarpeta='certificados',
                            usuario=request.user,
                        )
                        parametros_cliente.certificado_archivo = nuevo_archivo
                        parametros_cliente.save(
                            update_fields=['ambiente', 'clave_p12', 'certificado_archivo', 'actualizado_en']
                        )

                        if archivo_anterior and archivo_anterior.pk != nuevo_archivo.pk:
                            transaction.on_commit(
                                lambda archivo_anterior=archivo_anterior: eliminar_archivo(archivo_anterior)
                            )
                    else:
                        parametros_cliente.save(update_fields=['ambiente', 'clave_p12', 'actualizado_en'])

                    messages.success(request, 'Los parámetros de firma electrónica fueron actualizados.')

                elif accion == 'agregar_establecimiento':
                    codigo = request.POST.get('codigo', '').strip()
                    nombre = request.POST.get('nombre', '').strip()
                    direccion = request.POST.get('direccion', '').strip()
                    if len(codigo) != 3 or not codigo.isdigit():
                        raise ValueError('El código del establecimiento debe tener exactamente 3 dígitos.')
                    Establecimiento.objects.create(
                        cliente=cliente,
                        codigo=codigo,
                        nombre=nombre,
                        direccion=direccion,
                    )
                    messages.success(request, 'Establecimiento agregado correctamente.')

                elif accion == 'editar_establecimiento':
                    establecimiento = Establecimiento.objects.get(id=request.POST.get('id'), cliente=cliente)
                    codigo = request.POST.get('codigo', '').strip()
                    if len(codigo) != 3 or not codigo.isdigit():
                        raise ValueError('El código del establecimiento debe tener exactamente 3 dígitos.')
                    establecimiento.codigo = codigo
                    establecimiento.nombre = request.POST.get('nombre', '').strip()
                    establecimiento.direccion = request.POST.get('direccion', '').strip()
                    establecimiento.activo = request.POST.get('activo') == '1'
                    establecimiento.save(update_fields=['codigo', 'nombre', 'direccion', 'activo'])
                    messages.success(request, 'Establecimiento actualizado correctamente.')

                elif accion == 'agregar_punto':
                    establecimiento = Establecimiento.objects.get(id=request.POST.get('establecimiento_id'), cliente=cliente)
                    codigo = request.POST.get('codigo', '').strip()
                    if len(codigo) != 3 or not codigo.isdigit():
                        raise ValueError('El código del punto de emisión debe tener exactamente 3 dígitos.')
                    PuntoEmision.objects.create(
                        establecimiento=establecimiento,
                        codigo=codigo,
                        nombre=request.POST.get('nombre', '').strip(),
                    )
                    messages.success(request, 'Punto de emisión agregado correctamente.')

                elif accion == 'editar_punto':
                    punto = PuntoEmision.objects.select_related('establecimiento').get(
                        id=request.POST.get('id'),
                        establecimiento__cliente=cliente,
                    )
                    codigo = request.POST.get('codigo', '').strip()
                    if len(codigo) != 3 or not codigo.isdigit():
                        raise ValueError('El código del punto de emisión debe tener exactamente 3 dígitos.')
                    punto.codigo = codigo
                    punto.nombre = request.POST.get('nombre', '').strip()
                    punto.activo = request.POST.get('activo') == '1'
                    punto.save(update_fields=['codigo', 'nombre', 'activo'])
                    messages.success(request, 'Punto de emisión actualizado correctamente.')

                elif accion == 'guardar_secuencial':
                    punto = PuntoEmision.objects.select_related('establecimiento').get(
                        id=request.POST.get('punto_id'),
                        establecimiento__cliente=cliente,
                    )
                    tipo = request.POST.get('tipo_documento', '').strip()
                    secuencial = int(request.POST.get('secuencial_actual', '0'))
                    tipos_validos = dict(SecuencialDocumento.TipoDocumento.choices)
                    if tipo not in tipos_validos:
                        raise ValueError('El tipo de documento no es válido.')
                    if secuencial < 1:
                        raise ValueError('El secuencial debe ser mayor que cero.')
                    sec, creado = SecuencialDocumento.objects.get_or_create(
                        punto_emision=punto,
                        tipo_documento=tipo,
                        defaults={'secuencial_actual': secuencial},
                    )
                    if not creado:
                        sec.secuencial_actual = secuencial
                        sec.activo = request.POST.get('activo') == '1'
                        sec.save(update_fields=['secuencial_actual', 'activo'])
                    messages.success(request, 'Secuencial guardado correctamente.')

                else:
                    raise ValueError('La operación solicitada no es válida.')

        except (ValueError, TypeError, Establecimiento.DoesNotExist, PuntoEmision.DoesNotExist, SecuencialDocumento.DoesNotExist) as exc:
            messages.error(request, str(exc) or 'No fue posible completar la operación.')
        except IntegrityError:
            messages.error(request, 'No se pudo guardar porque el código o secuencial ya está registrado.')

        return redirect('parametros')

    establecimientos = cliente.establecimientos.prefetch_related('puntos_emision__secuenciales').all()
    return render(
        request,
        'parametros.html',
        {
            'cliente': cliente,
            'parametros_cliente': parametros_cliente,
            'establecimientos': establecimientos,
            'puede_editar': puede_editar,
            'tipos_documento': SecuencialDocumento.TipoDocumento.choices,
            'ambientes': ParametrosCliente.Ambiente.choices,
        },
    )


@login_required
def cambiar_contrasena(request):
    form = PasswordChangeForm(request.user, request.POST or None)

    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(
            request,
            'Tu contraseña fue actualizada correctamente.',
        )
        return redirect('cambiar_contrasena')

    context = {'form': form}

    if not request.user.is_staff and not request.user.is_superuser:
        asignacion = (
            UsuarioCliente.objects
            .filter(usuario=request.user, activo=True)
            .select_related('cliente')
            .first()
        )
        if asignacion is not None:
            context['cliente'] = asignacion.cliente

    return render(
        request,
        'cambiar-contrasena.html',
        context,
    )


def do_logout(request):
    logout(request)
    return redirect('home')


@login_required
def portal(request):
    if request.user.is_staff or request.user.is_superuser:
        return render(request, 'portal.html', {'es_administrador': True})

    asignaciones = UsuarioCliente.objects.filter(
        usuario=request.user,
        activo=True,
    ).select_related('cliente')

    for asignacion in asignaciones:
        permitido, _, _ = validar_acceso_cliente(asignacion.cliente)

        if permitido:
            request.session['cliente_ruccedcli'] = asignacion.cliente.ruccedcli
            resumen = obtener_resumen_cliente(asignacion.cliente)

            return render(
                request,
                'portal.html',
                {
                    'cliente': asignacion.cliente,
                    'resumen': resumen,
                    'rol': asignacion.rol,
                    'pendientes_legales': obtener_aceptaciones_pendientes(request.user),
                    'versiones_legales': versiones_legales_vigentes(),
                },
            )

    logout(request)
    messages.error(
        request,
        'El acceso al portal está restringido. Verifica el estado de tu cuenta.',
    )
    return redirect('signin')


@login_required
@require_POST
def aceptar_documento_legal(request):
    from .models import AceptacionLegal

    tipo = request.POST.get('tipo', '').strip()
    versiones = versiones_legales_vigentes()

    if tipo not in versiones:
        messages.error(request, 'El documento legal solicitado no es válido.')
        return redirect('portal')

    if tipo == AceptacionLegal.Tipo.POLITICA_DATOS:
        confirmado = request.POST.get('informado') == '1'
        if not confirmado:
            messages.error(request, 'Debes confirmar que has sido informado sobre la Política de Protección de Datos.')
            return redirect('portal')
    elif tipo == AceptacionLegal.Tipo.TERMINOS_USO:
        confirmado = request.POST.get('acepto') == '1'
        if not confirmado:
            messages.error(request, 'Debes aceptar los Términos y Condiciones de Uso para continuar.')

            return redirect('portal')

    version = versiones[tipo]
    AceptacionLegal.objects.get_or_create(
        usuario=request.user,
        tipo=tipo,
        version=version,
        defaults={
            'ip_address': request.META.get('REMOTE_ADDR'),
            'user_agent': request.META.get('HTTP_USER_AGENT', '')[:1000],
        },
    )
    return redirect('portal')


COMPRAS_COLUMNS = [
    ('numero', 'N°'),
    ('proveedor', 'PROVEEDOR'),
    ('ruc', 'RUC'),
    ('tipcom', 'TIP COM'),
    ('fecha', 'FECHA'),
    ('numfactura', 'NUMERO FACTURA'),
    ('numaut', 'NUM AUT.'),
    ('bases_sin_iva', 'BASES SIN IVA'),
    ('bases_con_iva', 'BASES CON IVA'),
    ('iva', 'IVA'),
    ('total', 'TOTAL'),
    ('retiva', 'RET IVA'),
    ('codret', 'COD RET'),
    ('retrenta', 'RET RENTA'),
    ('numret', 'NUM RET.'),
]


def _cliente_portal(request):
    if request.user.is_staff or request.user.is_superuser:
        return None

    asignacion = (
        UsuarioCliente.objects
        .filter(usuario=request.user, activo=True)
        .select_related('cliente')
        .first()
    )
    if asignacion is None:
        return None

    permitido, _, _ = validar_acceso_cliente(asignacion.cliente)
    return asignacion.cliente if permitido else None


def _cliente_db(cliente):
    """Obtiene una conexión PostgreSQL a la base de datos del cliente."""
    alias = f"cliente_{cliente.ruccedcli}"
    if alias not in connections.databases:
        base = settings.DATABASES['default'].copy()
        base['NAME'] = cliente.ruccedcli
        connections.databases[alias] = base
    return connections[alias]


def _compras_where(request):
    cliente = _cliente_portal(request)
    if cliente is None:
        return None, [], None

    where = []
    params = []

    hoy = datetime.now().date()
    primer_dia_mes = hoy.replace(day=1)

    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    proveedor = request.GET.get('proveedor', '').strip()

    if not fecha_desde:
        fecha_desde = primer_dia_mes.strftime('%Y-%m-%d')
    if not fecha_hasta:
        fecha_hasta = hoy.strftime('%Y-%m-%d')

    if fecha_desde:
        try:
            datetime.strptime(fecha_desde, '%Y-%m-%d')
            where.append("fecemi::date >= %s::date")
            params.append(fecha_desde)
        except ValueError:
            fecha_desde = ''

    if fecha_hasta:
        try:
            datetime.strptime(fecha_hasta, '%Y-%m-%d')
            where.append("fecemi::date < (%s::date + INTERVAL '1 day')")
            params.append(fecha_hasta)
        except ValueError:
            fecha_hasta = ''

    if proveedor:
        where.append("(ruccedprovee ILIKE %s OR nomprovee ILIKE %s)")
        params.extend([f'%{proveedor}%', f'%{proveedor}%'])

    return " AND ".join(where) if where else "1=1", params, {
        'cliente': cliente,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'proveedor': proveedor,
    }


def _compras_base_sql():
    return """
        SELECT
            c.fecemi,
            c.ruccedprovee,
            c.nomprovee,
            TRIM(c.tipcom::text) AS tipcom,
            c.numest,
            c.numptoemi,
            c.numsec,
            c.numaut,
            c.baseimpnoobj,
            c.baseimpiva0,
            c.baseexcenta,
            c.baseimpiva5,
            c.baseimpiva8,
            c.baseimpiva12,
            c.baseimpiva14,
            c.baseimpiva15,
            c.montoiva5,
            c.montoiva8,
            c.montoiva12,
            c.montoiva14,
            c.montoiva15,
            c.retencioniva10,
            c.retencioniva20,
            c.retencioniva30,
            c.retencioniva70,
            c.retencioniva100,
            c.codret,
            c.valret,
            c.numestret,
            c.numptoemiret,
            c.numsecret
        FROM comprasnue c
        WHERE TRIM(c.tipcom::text) IN ('01', '02')
    """


def _compras_query(where, params, cliente, limit=None, offset=None):
    sql = f"""
        SELECT
            ROW_NUMBER() OVER (ORDER BY fecemi::date ASC) AS numero,
            nomprovee AS proveedor,
            ruccedprovee AS ruc,
            CASE
                WHEN tipcom = '01' THEN 'FAC'
                WHEN tipcom = '02' THEN 'N/V'
                ELSE tipcom
            END AS tipcom,
            fecemi AS fecha,
            CONCAT(numest, '-', numptoemi, '-', numsec) AS numfactura,
            numaut,
            (
                COALESCE(NULLIF(baseimpnoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva0::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseexcenta::text, ''), '0')::numeric
            ) AS bases_sin_iva,
            (
                COALESCE(NULLIF(baseimpiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva15::text, ''), '0')::numeric
            ) AS bases_con_iva,
            (
                COALESCE(NULLIF(montoiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva15::text, ''), '0')::numeric
            ) AS iva,
            (
                COALESCE(NULLIF(baseimpnoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva0::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseexcenta::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva15::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva15::text, ''), '0')::numeric
            ) AS total,
            (
                COALESCE(NULLIF(retencioniva10::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva20::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva30::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva70::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva100::text, ''), '0')::numeric
            ) AS retiva,
            codret,
            COALESCE(NULLIF(valret::text, ''), '0')::numeric AS retrenta,
            CONCAT(numestret, '-', numptoemiret, '-', numsecret) AS numret
        FROM ({_compras_base_sql()}) compras_reporte
        WHERE {where}
        ORDER BY fecemi::date ASC
    """

    if limit is not None:
        sql += " LIMIT %s OFFSET %s"
        params = [*params, limit, offset or 0]

    with _cliente_db(cliente).cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.fetchall()


def _compras_resumen(where, params, cliente):
    sql = f"""
        SELECT
            COALESCE(SUM(
                COALESCE(NULLIF(baseimpnoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva0::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseexcenta::text, ''), '0')::numeric
            ), 0),
            COALESCE(SUM(
                COALESCE(NULLIF(baseimpiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva15::text, ''), '0')::numeric
            ), 0),
            COALESCE(SUM(
                COALESCE(NULLIF(montoiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva15::text, ''), '0')::numeric
            ), 0),
            COALESCE(SUM(
                COALESCE(NULLIF(baseimpnoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva0::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseexcenta::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseimpiva15::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva5::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva8::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva14::text, ''), '0')::numeric
                + COALESCE(NULLIF(montoiva15::text, ''), '0')::numeric
            ), 0),
            COALESCE(SUM(
                COALESCE(NULLIF(retencioniva10::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva20::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva30::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva70::text, ''), '0')::numeric
                + COALESCE(NULLIF(retencioniva100::text, ''), '0')::numeric
            ), 0),
            COALESCE(SUM(COALESCE(NULLIF(valret::text, ''), '0')::numeric), 0)
        FROM ({_compras_base_sql()}) compras_reporte
        WHERE {where}
    """

    with _cliente_db(cliente).cursor() as cursor:
        cursor.execute(sql, params)
        row = cursor.fetchone()

    return dict(zip(
        ['bases_sin_iva', 'bases_con_iva', 'iva', 'total', 'retiva', 'retrenta'],
        [float(value or 0) for value in row],
    ))


def _compras_datos_exportacion(request):
    where, params, filtros = _compras_where(request)
    if where is None:
        return None, [], None, None

    cliente = filtros['cliente']
    filas = _compras_query(where, params, cliente)
    resumen = _compras_resumen(where, params, cliente)
    return filtros, filas, resumen, where


@login_required
def compras(request):
    try:
        filtros, _, resumen, _ = _compras_datos_exportacion(request)
        if filtros is None:
            return redirect('portal')

        try:
            pagina = max(1, int(request.GET.get('pagina', '1')))
        except ValueError:
            pagina = 1

        por_pagina = 50
        where, params, _ = _compras_where(request)

        count_sql = f"SELECT COUNT(*) FROM ({_compras_base_sql()}) compras_reporte WHERE {where}"
        with _cliente_db(filtros['cliente']).cursor() as cursor:
            cursor.execute(count_sql, params)
            total_registros = cursor.fetchone()[0]

        offset = (pagina - 1) * por_pagina
        filas = _compras_query(where, params, filtros['cliente'], por_pagina, offset)
        total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)

        return render(request, 'compras.html', {
            'cliente': filtros['cliente'],
            'filas': filas,
            'resumen': resumen,
            'filtros': filtros,
            'pagina': pagina,
            'total_paginas': total_paginas,
            'total_registros': total_registros,
        })
    except Exception as exc:
        return HttpResponse(
            f"Error en reporte de compras: {type(exc).__name__}: {exc}",
            status=500,
            content_type='text/plain; charset=utf-8',
        )


@login_required
def compras_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    filtros, filas, resumen, _ = _compras_datos_exportacion(request)
    if filtros is None:
        return redirect('portal')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()

    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle

    estilo_titulo = ParagraphStyle(
        'ReporteComprasTitulo',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=16,
        alignment=TA_CENTER,
        spaceAfter=3,
    )
    estilo_cabecera = ParagraphStyle(
        'ReporteComprasCabecera',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=TA_CENTER,
        spaceAfter=2,
    )

    elements = [
        Paragraph('REPORTE DE COMPRAS', estilo_titulo),
        Paragraph(
            f"COMPRAS DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}",
            estilo_cabecera,
        ),
        Paragraph(
            f"{filtros['cliente'].nomclient} | RUC. {filtros['cliente'].ruccedcli}",
            estilo_cabecera,
        ),
        Spacer(1, 10),
    ]

    data = [[label for _, label in COMPRAS_COLUMNS]]
    for row in filas:
        formatted = []
        for index, value in enumerate(row):
            if index in (0,):
                formatted.append(str(value))
            elif index in (1, 2, 3, 4, 5, 6, 13, 14):
                formatted.append(str(value or ''))
            else:
                formatted.append(f"{float(value or 0):.2f}")
        data.append(formatted)

    data.append([
        '', '', '', '', '', '', '',
        f"{resumen['bases_sin_iva']:.2f}",
        f"{resumen['bases_con_iva']:.2f}",
        f"{resumen['iva']:.2f}",
        f"{resumen['total']:.2f}",
        f"{resumen['retiva']:.2f}",
        '',
        f"{resumen['retrenta']:.2f}",
        '',
    ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[24, 105, 75, 40, 52, 78, 55, 58, 58, 45, 58, 48, 42, 55, 60],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (0, -1), 'CENTER'),
        ('ALIGN', (7, 1), (14, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_compras.pdf"'
    return response


@login_required
def compras_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    filtros, filas, resumen, _ = _compras_datos_exportacion(request)
    if filtros is None:
        return redirect('portal')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compras'
    ws.append([label for _, label in COMPRAS_COLUMNS])

    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')

    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', 'RESUMEN'])
    ws.cell(ws.max_row, 8, resumen['bases_sin_iva'])
    ws.cell(ws.max_row, 9, resumen['bases_con_iva'])
    ws.cell(ws.max_row, 10, resumen['iva'])
    ws.cell(ws.max_row, 11, resumen['total'])
    ws.cell(ws.max_row, 12, resumen['retiva'])
    ws.cell(ws.max_row, 14, resumen['retrenta'])

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    widths = [7, 38, 17, 10, 13, 25, 18, 16, 16, 13, 16, 14, 11, 16, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_compras.xlsx"'
    return response

def about(request):
    return render(request, 'about.html')


def contactanos(request):
    return render(request, 'contactanos.html')


def avisos_legales(request):
    return render(request, 'avisos-legales.html')


def servicios(request):
    raise Http404


def blog(request):
    raise Http404


def error_404(request, exception):
    return render(request, '404.html', status=404)


@staff_member_required
def template_catalog(request):
    """Catálogo interno para revisar las plantillas del tema antes de reutilizarlas."""
    return render(request, 'template-catalog.html', {'templates': TEMPLATE_PREVIEWS})


@staff_member_required
def template_preview(request, slug):
    """Renderiza una plantilla del catálogo usando el mismo sistema de templates de Django."""
    template_name = TEMPLATE_PREVIEWS.get(slug)
    if template_name is None:
        raise Http404
    return render(request, template_name, {'preview_mode': True, 'preview_slug': slug})
