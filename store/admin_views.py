from io import BytesIO
from functools import wraps
import mimetypes
from pathlib import Path

from django.contrib import messages
from django.db import connection, connections
from django.db.models import Count
from django.db.models.functions import TruncDate
from django.utils import timezone
from datetime import datetime, timedelta
from urllib.parse import urlencode
from django.shortcuts import get_object_or_404, redirect, render
from django.conf import settings
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.hashers import check_password, make_password

from .models import AdminPerfil, Cliente, UsuarioCliente, VisitaWeb, Suscriptor


ADMIN_SESSION_KEY = 'admin_portal'
ADMIN_USERNAME_KEY = 'admin_username'
ADMIN_NAME_KEY = 'admin_name'
ADMIN_WORKER_KEY = 'admin_worker'


def admin_required(view_func):
    """Protege las vistas del portal administrativo con la sesión propia de TotalCounts."""
    @wraps(view_func)
    def wrapped(request, *args, **kwargs):
        if not request.session.get(ADMIN_SESSION_KEY):
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return wrapped


def _admin_required(request):
    return request.session.get(ADMIN_SESSION_KEY, False)


def admin_login(request):
    if request.session.get(ADMIN_SESSION_KEY):
        return redirect('admin_dashboard')

    error = ''
    username = ''

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')

        if not username or not password:
            error = 'Ingresa el usuario y la contraseña.'
        else:
            with connection.cursor() as cursor:
                cursor.execute(
                    '''
                    SELECT usrname, nomusuari, conusuari
                    FROM seguridad
                    WHERE UPPER(TRIM(usrname::text)) = UPPER(TRIM(%s))
                    LIMIT 1
                    ''',
                    [username],
                )
                usuario = cursor.fetchone()

            if usuario:
                usuario_db, nombre_db, clave_db = usuario
                clave_db = '' if clave_db is None else str(clave_db)
                valido = False

                # Soporta tanto las claves heredadas almacenadas como texto
                # como hashes Django, si en el futuro se migran.
                if clave_db.startswith(('pbkdf2_', 'argon2$', 'bcrypt$', 'scrypt$')):
                    try:
                        valido = check_password(password, clave_db)
                    except Exception:
                        valido = False
                else:
                    valido = password == clave_db

                if valido:
                    request.session[ADMIN_SESSION_KEY] = True
                    request.session[ADMIN_USERNAME_KEY] = str(usuario_db).strip()
                    request.session[ADMIN_NAME_KEY] = str(nombre_db or usuario_db).strip()
                    request.session[ADMIN_WORKER_KEY] = False
                    request.session.set_expiry(1800)
                    return redirect('admin_dashboard')

            error = 'Usuario o contraseña inválidos.'

    return render(
        request,
        'admin/login.html',
        {
            'error': error,
            'username': username,
        },
    )



@csrf_exempt
def conta_worker_error(request):
    """Recibe errores internos de Conta y los reporta por el SMTP de TtCWeb."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)

    token = request.headers.get('X-TotalCounts-Internal', '').strip()
    if not token or token != settings.CONTA_INTERNAL_TOKEN:
        return JsonResponse({'error': 'No autorizado.'}, status=401)

    try:
        import json
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except (UnicodeDecodeError, json.JSONDecodeError):
        return JsonResponse({'error': 'JSON inválido.'}, status=400)

    worker = str(payload.get('worker') or '-').strip()
    usuario = str(payload.get('usuario') or '-').strip()
    equipo = str(payload.get('equipo') or '-').strip()
    job_id = str(payload.get('job_id') or '-').strip()
    ruc = str(payload.get('ruc') or '-').strip()
    anio = str(payload.get('anio') or '-').strip()
    mes = str(payload.get('mes') or '-').strip()
    tipo = str(payload.get('tipo_comprobante') or '-').strip()
    operacion = str(payload.get('operacion') or '-').strip()
    error_tipo = str(payload.get('error_tipo') or 'Exception').strip()
    error_mensaje = str(payload.get('error_mensaje') or '-').strip()
    traceback_texto = str(payload.get('traceback') or '-')

    asunto = f'[Conta] Error SRI Worker | {worker} | job {job_id}'
    cuerpo = f"""Conta - ERROR AUTOMÁTICO DEL SRI WORKER

Worker: {worker}
Usuario: {usuario}
Equipo: {equipo}
Job ID: {job_id}
RUC: {ruc}
Año: {anio}
Mes: {mes}
Tipo comprobante: {tipo}
Operación: {operacion}

Error: {error_tipo}: {error_mensaje}

Traceback completo:
{traceback_texto}
"""

    try:
        from django.core.mail import send_mail
        send_mail(
            asunto,
            cuerpo,
            settings.DEFAULT_FROM_EMAIL,
            [getattr(settings, 'BUG_REPORT_RECIPIENT', 'eldudyaguirre@gmail.com')],
            fail_silently=False,
        )
        return JsonResponse({'ok': True})
    except Exception as exc:
        return JsonResponse({'error': f'No se pudo enviar el correo: {exc}'}, status=500)


def admin_logout(request):
    request.session.pop(ADMIN_SESSION_KEY, None)
    request.session.pop(ADMIN_USERNAME_KEY, None)
    request.session.pop(ADMIN_NAME_KEY, None)
    request.session.pop(ADMIN_WORKER_KEY, None)
    return redirect('admin_login')


def _admin_conciliacion_cliente(cliente, anio=None):
    """Obtiene la conciliación mensual de ventas y compras."""
    db_name = str(cliente.ruccedcli).strip()
    alias = f'cliente_{db_name}'
    if alias not in connections.databases:
        base = settings.DATABASES['default'].copy()
        base['NAME'] = db_name
        connections.databases[alias] = base

    db = connections[alias]

    with db.cursor() as cursor:
        cursor.execute("""
            SELECT DISTINCT anio
            FROM (
                SELECT EXTRACT(YEAR FROM fecfactur::date)::integer AS anio
                FROM ventas
                WHERE fecfactur IS NOT NULL
                UNION
                SELECT TRIM("año"::text)::integer AS anio
                FROM comprasnue
                WHERE TRIM("año"::text) ~ '^\d{4}$'
            ) periodos
            WHERE anio IS NOT NULL
            ORDER BY anio DESC
        """)
        anios = [int(row[0]) for row in cursor.fetchall()]

    if not anios:
        return {
            'anio': anio or 0, 'anios': [], 'meses': [],
            'mes_final_nombre': '',
            'totales': {
                'ventas': 0, 'compras': 0, 'retrenta': 0,
                'retiva': 0, 'resultado': 0, 'retenciones': 0,
            },
        }

    try:
        anio = int(anio)
    except (TypeError, ValueError):
        anio = anios[0]

    if anio not in anios:
        anio = anios[0]

    with db.cursor() as cursor:
        cursor.execute("""
            WITH meses AS (
                SELECT generate_series(1, 12) AS mes
            ),
            ventas_mes AS (
                SELECT
                    EXTRACT(MONTH FROM fecfactur::date)::integer AS mes,
                    COALESCE(SUM(
                        COALESCE(NULLIF(NULLIF(TRIM(basenoobj::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva0::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva5::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva8::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva12::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva14::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseiva15::text), ''), '.'), '0')::numeric
                    ), 0) AS ventas,
                    COALESCE(SUM(COALESCE(NULLIF(NULLIF(TRIM(retrenta::text), ''), '.'), '0')::numeric), 0) AS retrenta,
                    COALESCE(SUM(COALESCE(NULLIF(NULLIF(TRIM(retiva::text), ''), '.'), '0')::numeric), 0) AS retiva
                FROM ventas
                WHERE fecfactur IS NOT NULL
                  AND EXTRACT(YEAR FROM fecfactur::date)::integer = %s
                GROUP BY EXTRACT(MONTH FROM fecfactur::date)::integer
            ),
            compras_mes AS (
                SELECT
                    TRIM(mes::text)::integer AS mes,
                    COALESCE(SUM(
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpnoobj::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva0::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseexenta::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva5::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva8::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva12::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva14::text), ''), '.'), '0')::numeric +
                        COALESCE(NULLIF(NULLIF(TRIM(baseimpiva15::text), ''), '.'), '0')::numeric
                    ), 0) AS compras
                FROM comprasnue
                WHERE TRIM("año"::text) = %s
                  AND TRIM(mes::text) ~ '^\d{1,2}$'
                  AND TRIM(mes::text)::integer BETWEEN 1 AND 12
                  AND TRIM(tipcom::text) IN ('01', '02')
                GROUP BY TRIM(mes::text)::integer
            )
            SELECT
                m.mes,
                COALESCE(v.ventas, 0),
                COALESCE(c.compras, 0),
                COALESCE(v.retrenta, 0),
                COALESCE(v.retiva, 0)
            FROM meses m
            LEFT JOIN ventas_mes v ON v.mes = m.mes
            LEFT JOIN compras_mes c ON c.mes = m.mes
            ORDER BY m.mes
        """, [str(anio), str(anio)])
        rows = cursor.fetchall()

    nombres = [
        'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO',
        'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE'
    ]

    meses = []
    for mes, ventas, compras, retrenta, retiva in rows:
        meses.append({
            'nombre': nombres[int(mes) - 1],
            'ventas': ventas or 0,
            'compras': compras or 0,
            'retrenta': retrenta or 0,
            'retiva': retiva or 0,
        })

    totales = {
        campo: sum((m[campo] for m in meses), 0)
        for campo in ('ventas', 'compras', 'retrenta', 'retiva')
    }
    totales['resultado'] = totales['ventas'] - totales['compras']
    totales['retenciones'] = totales['retrenta'] + totales['retiva']

    return {
        'anio': anio,
        'anios': anios,
        'meses': meses,
        'mes_final_nombre': 'DICIEMBRE',
        'totales': totales,
    }


@admin_required
def admin_workers(request):
    """Administra qué usuarios administrativos tienen habilitado el worker SRI."""
    if request.method == 'POST':
        usuario = request.POST.get('usuario', '').strip()
        accion = request.POST.get('accion', '').strip()

        if usuario and accion in ('activar', 'desactivar'):
            valor = accion == 'activar'
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE seguridad
                    SET worker = %s
                    WHERE UPPER(TRIM(usrname::text)) = UPPER(TRIM(%s))
                    """,
                    [valor, usuario],
                )
            messages.success(
                request,
                f'Worker {"activado" if valor else "desactivado"} para {usuario}.'
            )

        return redirect('admin_workers')

    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT
                TRIM(usrname::text) AS usuario,
                COALESCE(TRIM(nomusuari::text), '') AS nombre,
                COALESCE(worker, FALSE) AS worker
            FROM seguridad
            ORDER BY worker DESC, nomusuari, usrname
            """
        )
        trabajadores = [
            {
                'usuario': row[0],
                'nombre': row[1],
                'worker': bool(row[2]),
            }
            for row in cursor.fetchall()
        ]

    return render(
        request,
        'admin/workers.html',
        {
            'trabajadores': trabajadores,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_datos_usuario(request):
    """Muestra y permite editar los datos del usuario administrativo autenticado."""
    usuario = request.session.get(ADMIN_USERNAME_KEY, '').strip()
    nombre_sesion = request.session.get(ADMIN_NAME_KEY, '').strip()

    perfil, _ = AdminPerfil.objects.get_or_create(
        usuario=usuario,
        defaults={'nombres': nombre_sesion},
    )

    if request.method == 'POST':
        perfil.nombres = request.POST.get('nombres', '').strip()
        perfil.direccion = request.POST.get('direccion', '').strip()
        perfil.telefono = request.POST.get('telefono', '').strip()
        perfil.email = request.POST.get('email', '').strip()

        fecha_nacimiento = request.POST.get('fecha_nacimiento', '').strip()
        perfil.fecha_nacimiento = fecha_nacimiento or None

        nueva_foto = request.FILES.get('foto')
        if nueva_foto:
            perfil.foto = nueva_foto

        perfil.save()

        # Mantiene el nombre mostrado en la barra superior sincronizado
        # con el nombre guardado en el perfil.
        request.session[ADMIN_NAME_KEY] = perfil.nombres or usuario

        messages.success(request, 'Datos de usuario actualizados correctamente.')
        return redirect('admin_datos_usuario')

    return render(
        request,
        'admin/datos_usuario.html',
        {
            'admin_usuario': usuario,
            'admin_nombre': perfil.nombres or nombre_sesion,
            'perfil': perfil,
        },
    )


@admin_required
def admin_cambiar_contrasena(request):
    """Permite cambiar la contraseña del usuario administrativo en la tabla legacy seguridad."""
    usuario_sesion = request.session.get(ADMIN_USERNAME_KEY, '').strip()
    error = ''

    if request.method == 'POST':
        actual = request.POST.get('old_password', '')
        nueva = request.POST.get('new_password1', '')
        confirmacion = request.POST.get('new_password2', '')

        if not actual or not nueva or not confirmacion:
            error = 'Completa todos los campos.'
        elif len(nueva) < 8:
            error = 'La nueva contraseña debe tener al menos 8 caracteres.'
        elif nueva != confirmacion:
            error = 'La confirmación de la nueva contraseña no coincide.'
        else:
            with connection.cursor() as cursor:
                cursor.execute(
                    '''
                    SELECT conusuari
                    FROM seguridad
                    WHERE UPPER(TRIM(usrname::text)) = UPPER(TRIM(%s))
                    LIMIT 1
                    ''',
                    [usuario_sesion],
                )
                fila = cursor.fetchone()

            if not fila:
                error = 'No se encontró el usuario administrativo.'
            else:
                clave_actual = '' if fila[0] is None else str(fila[0])
                es_hash = clave_actual.startswith(
                    ('pbkdf2_', 'argon2', 'bcrypt', 'scrypt')
                )

                if es_hash:
                    try:
                        valido = check_password(actual, clave_actual)
                    except Exception:
                        valido = False
                else:
                    valido = actual == clave_actual

                if not valido:
                    error = 'La contraseña actual no es correcta.'
                elif nueva == actual:
                    error = 'La nueva contraseña debe ser diferente a la actual.'
                else:
                    nueva_guardada = make_password(nueva) if es_hash else nueva

                    with connection.cursor() as cursor:
                        cursor.execute(
                            '''
                            UPDATE seguridad
                            SET conusuari = %s
                            WHERE UPPER(TRIM(usrname::text)) = UPPER(TRIM(%s))
                            ''',
                            [nueva_guardada, usuario_sesion],
                        )

                    messages.success(
                        request,
                        'Tu contraseña fue actualizada correctamente.',
                    )
                    return redirect('admin_cambiar_contrasena')


    return render(
        request,
        'admin/cambiar_contrasena.html',
        {
            'admin_usuario': usuario_sesion,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'password_error': error,
        },
    )


@admin_required
def admin_foto_usuario(request):
    """Entrega la foto del perfil administrativo autenticado."""
    usuario = request.session.get(ADMIN_USERNAME_KEY, '').strip()
    perfil = AdminPerfil.objects.filter(usuario=usuario).first()

    if not perfil or not perfil.foto:
        raise Http404

    try:
        return FileResponse(
            perfil.foto.open('rb'),
            content_type=mimetypes.guess_type(perfil.foto.name)[0] or 'application/octet-stream',
        )
    except FileNotFoundError:
        raise Http404


def _admin_dashboard_saldos():
    """Obtiene los saldos globales que se muestran en el panel administrativo."""
    cuentas_por_cobrar = (
        Cliente.objects.aggregate(
            total=__import__('django.db.models', fromlist=['Sum']).Sum('salcuenta')
        )['total'] or 0
    )

    # Los valores pendientes de facturar se encuentran en la tabla
    # prefactura de la base maestra y se calculan sumando preitefac.
    saldo_por_facturar = 0
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COALESCE(SUM(preitefac), 0) FROM prefactura"
            )
            saldo_por_facturar = cursor.fetchone()[0] or 0
    except Exception:
        saldo_por_facturar = 0

    return {
        'cuentas_por_cobrar': cuentas_por_cobrar,
        'saldo_por_facturar': saldo_por_facturar,
    }

def admin_dashboard(request):
    clientes_qs = Cliente.objects.all()
    clientes = clientes_qs.order_by('nomclient')[:12]

    hoy = timezone.localdate()
    inicio_mes = hoy.replace(day=1)
    inicio_mes_anterior = (inicio_mes - timedelta(days=1)).replace(day=1)

    visitas_hoy_qs = VisitaWeb.objects.filter(fecha_hora__date=hoy)
    visitas_mes_qs = VisitaWeb.objects.filter(fecha_hora__date__gte=inicio_mes)
    visitas_mes_anterior_qs = VisitaWeb.objects.filter(
        fecha_hora__date__gte=inicio_mes_anterior,
        fecha_hora__date__lt=inicio_mes,
    )

    inicio_grafica = hoy - timedelta(days=6)
    visitas_grafica = (
        VisitaWeb.objects
        .filter(fecha_hora__date__gte=inicio_grafica, fecha_hora__date__lte=hoy)
        .annotate(dia=TruncDate('fecha_hora'))
        .values('dia')
        .annotate(total=Count('id'))
        .order_by('dia')
    )
    visitas_por_dia = {item['dia']: item['total'] for item in visitas_grafica}

    grafica_visitas = [
        {
            'fecha': inicio_grafica + timedelta(days=i),
            'total': visitas_por_dia.get(inicio_grafica + timedelta(days=i), 0),
        }
        for i in range(7)
    ]

    # Solo mostramos páginas públicas reales del sitio. Las rutas técnicas,
    # administrativas, del portal y solicitudes de scanners quedan fuera.
    paginas_publicas = (
        '/',
        '/about/',
        '/contactanos/',
        '/avisos-legales/',
        '/servicios/',
        '/blog/',
    )
    paginas_mas_visitadas = (
        visitas_mes_qs
        .filter(pagina__in=paginas_publicas)
        .values('pagina')
        .annotate(total=Count('id'))
        .order_by('-total')[:8]
    )

    saldos = _admin_dashboard_saldos()

    context = {
        'clientes_total': clientes_qs.count(),
        'clientes_activos': clientes_qs.filter(activo=True).count(),
        'clientes_inactivos': clientes_qs.filter(activo=False).count(),
        'usuarios_clientes': UsuarioCliente.objects.filter(activo=True).count(),
        'cuentas_por_cobrar': saldos['cuentas_por_cobrar'],
        'saldo_por_facturar': saldos['saldo_por_facturar'],
        'clientes': clientes,
        'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
        'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        'visitas_hoy': visitas_hoy_qs.count(),
        'visitantes_hoy': visitas_hoy_qs.values('visitor_id').distinct().count(),
        'visitas_mes': visitas_mes_qs.count(),
        'visitas_mes_anterior': visitas_mes_anterior_qs.count(),
        'grafica_visitas': grafica_visitas,
        'paginas_mas_visitadas': paginas_mas_visitadas,
    }
    return render(request, 'admin/dashboard.html', context)


@admin_required
def admin_contrasenas_editar(request, ruc):
    """Edita únicamente las credenciales del cliente desde el módulo Contraseñas."""
    cliente = get_object_or_404(Cliente, pk=ruc)

    if request.method != 'POST':
        return redirect('admin_contrasenas')

    campos = (
        'clavesri',
        'cediess',
        'claveiess',
        'mrlcon',
        'mrlsal',
        'clavesuper',
        'iessdomestica',
    )

    for campo in campos:
        setattr(cliente, campo, request.POST.get(campo, '').strip())

    tipdec = request.POST.get('tipdec', '').strip().upper()
    if tipdec in ('MENSUAL', 'SEMESTRAL', 'ANUAL'):
        # El campo existente en el modelo es semensual (TipDec).
        cliente.semensual = tipdec

    campos_guardar = list(campos)
    if tipdec in ('MENSUAL', 'SEMESTRAL', 'ANUAL'):
        campos_guardar.append('semensual')

    cliente.save(update_fields=campos_guardar)

    messages.success(
        request,
        f'Las contraseñas de {cliente.nomclient} fueron actualizadas correctamente.',
    )

    query = request.POST.get('q', '').strip()
    url = redirect('admin_contrasenas')
    if query:
        url['Location'] += '?' + urlencode({'q': query})
    return url


@admin_required
def admin_contrasenas_toggle_activo(request, ruc):
    """Activa o desactiva un cliente desde el módulo Contraseñas."""
    cliente = get_object_or_404(Cliente, pk=ruc)

    if request.method != 'POST':
        return redirect('admin_contrasenas')

    cliente.activo = not bool(cliente.activo)
    cliente.save(update_fields=['activo'])

    estado = 'activado' if cliente.activo else 'desactivado'
    messages.success(request, f'El cliente {cliente.nomclient} fue {estado} correctamente.')

    query = request.POST.get('q', '').strip()
    url = redirect('admin_contrasenas')
    if query:
        url['Location'] += f'?q={query}'
    return url


@admin_required
def admin_clientes(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    clientes = Cliente.objects.all().order_by('nomclient')

    if query:
        from django.db.models import Q
        clientes = clientes.filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if estado == 'activos':
        clientes = clientes.filter(activo=True)
    elif estado == 'inactivos':
        clientes = clientes.filter(activo=False)

    if dia != 'todos':
        clientes = clientes.filter(diadeclaracion=dia)

    if tipdec:
        clientes = clientes.filter(tipdec__iexact=tipdec)

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/clientes.html',
        {
            'clientes': clientes,
            'query': query,
            'estado': estado,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
        },
    )


@admin_required
def admin_suscriptores(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()

    suscriptores = Suscriptor.objects.all()

    if query:
        from django.db.models import Q
        suscriptores = suscriptores.filter(
            Q(nombre__icontains=query) |
            Q(email__icontains=query)
        )

    if estado == 'activos':
        suscriptores = suscriptores.filter(activo=True)
    elif estado == 'inactivos':
        suscriptores = suscriptores.filter(activo=False)

    return render(
        request,
        'admin/suscriptores.html',
        {
            'suscriptores': suscriptores,
            'query': query,
            'estado': estado,
            'suscriptores_total': Suscriptor.objects.count(),
            'suscriptores_activos': Suscriptor.objects.filter(activo=True).count(),
        },
    )


@admin_required
def admin_contrasenas(request):
    """Listado administrativo de clientes activos y sus credenciales tributarias/laborales."""
    query = request.GET.get('q', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    # Al entrar se muestran únicamente clientes activos.
    # Cuando se realiza una búsqueda, se consulta toda la base de datos,
    # incluyendo clientes inactivos.
    clientes_qs = Cliente.objects.filter(activo=True)

    if query:
        from django.db.models import Q
        clientes_qs = Cliente.objects.all().filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if dia != 'todos':
        clientes_qs = clientes_qs.filter(diadeclaracion=str(dia))

    if tipdec:
        clientes_qs = clientes_qs.filter(semensual__iexact=tipdec)

    clientes_qs = clientes_qs.order_by('nomclient')

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/contrasenas.html',
        {
            'clientes': clientes_qs,
            'pestanas': [
                {
                    'dia': dia_item,
                    'clientes': clientes_qs.filter(diadeclaracion=str(dia_item)),
                }
                for dia_item in dias
            ],
            'clientes_busqueda': clientes_qs if query else Cliente.objects.none(),
            'query': query,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


def _admin_compras_filtros(request):
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)
    fecha_desde = request.GET.get('fecha_desde', '').strip() or primer_dia_mes.strftime('%Y-%m-%d')
    fecha_hasta = request.GET.get('fecha_hasta', '').strip() or hoy.strftime('%Y-%m-%d')
    proveedor = request.GET.get('proveedor', '').strip()

    where = []
    params = []
    try:
        datetime.strptime(fecha_desde, '%Y-%m-%d')
        where.append('fecemi::date >= %s::date')
        params.append(fecha_desde)
    except ValueError:
        fecha_desde = ''

    try:
        datetime.strptime(fecha_hasta, '%Y-%m-%d')
        where.append("fecemi::date < (%s::date + INTERVAL '1 day')")
        params.append(fecha_hasta)
    except ValueError:
        fecha_hasta = ''

    if proveedor:
        where.append('(ruccedprovee ILIKE %s OR nomprovee ILIKE %s)')
        params.extend([f'%{proveedor}%', f'%{proveedor}%'])

    return (' AND '.join(where) if where else '1=1'), params, {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'proveedor': proveedor,
    }


def _admin_ventas_query(where, params, cliente, limit=None, offset=None):
    """
    Usa exactamente la misma consulta de ventas que el portal.

    La tabla ventas de este sistema almacena la base gravada en baseiva12
    y el IVA total en iva. No debemos reconstruirlos sumando columnas iva5,
    iva8, iva12, etc., porque esas columnas no son la fuente usada por este
    reporte y pueden venir vacías.
    """
    from .views import _ventas_query
    return _ventas_query(where, params, cliente, limit, offset)


def _admin_ventas_resumen(where, params, cliente):
    """
    Usa exactamente el mismo resumen tributario que el reporte de ventas
    del portal.
    """
    from .views import _ventas_resumen
    return _ventas_resumen(where, params, cliente)

def _admin_ventas_filtros(request):
    hoy = timezone.localdate()
    primer_dia = hoy.replace(day=1)
    desde = request.GET.get('fecha_desde','').strip() or primer_dia.strftime('%Y-%m-%d')
    hasta = request.GET.get('fecha_hasta','').strip() or hoy.strftime('%Y-%m-%d')
    cliente_busqueda = request.GET.get('cliente_busqueda','').strip()
    where=[]; params=[]
    try:
        datetime.strptime(desde,'%Y-%m-%d'); where.append('fecfactur::date >= %s::date'); params.append(desde)
    except ValueError: desde=''
    try:
        datetime.strptime(hasta,'%Y-%m-%d'); where.append("fecfactur::date < (%s::date + INTERVAL '1 day')"); params.append(hasta)
    except ValueError: hasta=''
    if cliente_busqueda:
        where.append('(ruccedcli ILIKE %s OR nomcli ILIKE %s)')
        params.extend([f'%{cliente_busqueda}%',f'%{cliente_busqueda}%'])
    return (' AND '.join(where) if where else '1=1'), params, {'fecha_desde':desde,'fecha_hasta':hasta,'cliente_busqueda':cliente_busqueda}


@admin_required
def admin_ventas_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente=get_object_or_404(Cliente,pk=ruc); where,params,filtros=_admin_ventas_filtros(request)
    filas=_admin_ventas_query(where,params.copy(),cliente); resumen=_admin_ventas_resumen(where,params.copy(),cliente)
    wb=Workbook(); ws=wb.active; ws.title='Ventas'
    ws.append([x[1] for x in VENTAS_COLUMNS])
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='21333E'); cell.alignment=Alignment(horizontal='center')
    for row in filas: ws.append(list(row))
    ws.append([]); ws.append(['','','','','','RESUMEN',resumen['base0'],resumen['baseiva'],resumen['iva'],resumen['total'],resumen['retiva'],resumen['retrenta']])
    buffer=BytesIO(); wb.save(buffer)
    response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']='attachment; filename="ventas_administrativo.xlsx"'; return response


@admin_required
def admin_ventas_pdf(request, ruc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape,A4
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
    from .views import _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente=get_object_or_404(Cliente,pk=ruc); where,params,filtros=_admin_ventas_filtros(request)
    filas=_ventas_query(where,params.copy(),cliente); resumen=_ventas_resumen(where,params.copy(),cliente)
    buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=20,rightMargin=20,topMargin=20,bottomMargin=20)
    styles=getSampleStyleSheet(); title=ParagraphStyle('avt',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=14,alignment=TA_CENTER)
    head=ParagraphStyle('avh',parent=styles['Normal'],fontName='Helvetica-Bold',fontSize=8,alignment=TA_CENTER)
    cell=ParagraphStyle('avc',parent=styles['Normal'],fontName='Helvetica',fontSize=5.4,leading=6,alignment=TA_CENTER,wordWrap='CJK')
    # Un solo estilo para TODAS las celdas: texto y valores numéricos.
    def celda(valor='', estilo=cell):
        return Paragraph(str(valor if valor is not None else ''), estilo)

    data=[[celda(x[1], head) for x in VENTAS_COLUMNS]]
    for row in filas:
        values=[]
        for i,v in enumerate(row):
            if i in (0,1,2,3,4,5,12,13):
                texto = str(v or '')
            else:
                try:
                    texto = f'{float(v or 0):.2f}'
                except (TypeError,ValueError):
                    texto = str(v or '')
            values.append(celda(texto))
        data.append(values)
    data.append([
        celda(''), celda(''), celda(''), celda(''), celda(''), celda(''),
        celda(f"{resumen['base0']:.2f}"),
        celda(f"{resumen['baseiva']:.2f}"),
        celda(f"{resumen['iva']:.2f}"),
        celda(f"{resumen['total']:.2f}"),
        celda(f"{resumen['retiva']:.2f}"),
        celda(f"{resumen['retrenta']:.2f}"),
        celda(''), celda('')
    ])
    table=Table(data,repeatRows=1,colWidths=[22,105,68,48,72,76,52,52,42,52,45,50,52,60])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#21333e')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#d8e0e3')),('ALIGN',(0,0),(-1,-1),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#eef5f5'))]))
    doc.build([Paragraph('REPORTE DE VENTAS',title),Paragraph(f"VENTAS DESDE {filtros['fecha_desde']} A {filtros['fecha_hasta']}",head),Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}",head),Spacer(1,10),table])
    response=HttpResponse(buf.getvalue(),content_type='application/pdf'); response['Content-Disposition']='attachment; filename="ventas_administrativo_v2.pdf"'; response['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0'; response['Pragma']='no-cache'; response['Expires']='0'; return response


@admin_required
def admin_ventas_editar(request, ruc):
    from .views import _cliente_db
    cliente=get_object_or_404(Cliente,pk=ruc)
    numfactur=(request.POST.get('numfactur') if request.method=='POST' else request.GET.get('numfactur','')).strip()
    db=_cliente_db(cliente)
    try:
        with db.cursor() as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_schema='public' AND table_name='ventas' ORDER BY ordinal_position")
            cols=[x[0] for x in cur.fetchall()]
            if 'numfactur' not in cols: return JsonResponse({'ok':False,'error':'El campo numfactur no existe en ventas.'},status=500)
            wanted=['fecfactur','numfactur','autorizacion','basenoobj','baseiva0','baseiva5','baseiva8','baseiva12','baseiva14','baseiva15','iva','iva5','iva8','iva12','iva14','iva15','retiva','retrenta','numret','autret','codret']
            select=[x for x in wanted if x in cols]
            cur.execute(f'SELECT {",".join(chr(34)+x+chr(34) for x in select)} FROM ventas WHERE "numfactur"::text=%s LIMIT 1',[numfactur])
            row=cur.fetchone()
            if not row: return JsonResponse({'ok':False,'error':f'No se encontró la factura {numfactur}.'},status=404)
            data=dict(zip(select,row))
            if request.method=='GET':
                for k,v in data.items(): data[k]='' if v is None else (v.isoformat() if hasattr(v,'isoformat') else str(v))
                return JsonResponse({'ok':True,'venta':data})
            editable=[x for x in ['fecfactur','autorizacion','basenoobj','baseiva0','baseiva5','baseiva8','baseiva12','baseiva14','baseiva15','iva5','iva8','iva12','iva14','iva15','retiva','retrenta','numret','autret','codret'] if x in cols and x in request.POST]
            vals={x:(request.POST.get(x,'').strip() or None) for x in editable}
            if not vals: return JsonResponse({'ok':False,'error':'No hay cambios para guardar.'},status=400)
            sets=', '.join(f'"{x}"=%s' for x in vals)
            cur.execute(f'UPDATE ventas SET {sets} WHERE "numfactur"::text=%s',[*vals.values(),numfactur])
            db.commit()
            return JsonResponse({'ok':True})
    except Exception as exc:
        try: db.rollback()
        except Exception: pass
        return JsonResponse({'ok':False,'error':f'Error actualizando venta: {type(exc).__name__}: {exc}'},status=500)



def _admin_notas_credito_filtros(request):
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)
    fecha_desde = request.GET.get('fecha_desde', '').strip() or primer_dia_mes.strftime('%Y-%m-%d')
    fecha_hasta = request.GET.get('fecha_hasta', '').strip() or hoy.strftime('%Y-%m-%d')
    proveedor = request.GET.get('proveedor', '').strip()

    where = []
    params = []
    try:
        datetime.strptime(fecha_desde, '%Y-%m-%d')
        where.append('fecemi::date >= %s::date')
        params.append(fecha_desde)
    except ValueError:
        fecha_desde = ''

    try:
        datetime.strptime(fecha_hasta, '%Y-%m-%d')
        where.append("fecemi::date < (%s::date + INTERVAL '1 day')")
        params.append(fecha_hasta)
    except ValueError:
        fecha_hasta = ''

    if proveedor:
        where.append('(ruccedprovee ILIKE %s OR nomprovee ILIKE %s)')
        params.extend([f'%{proveedor}%', f'%{proveedor}%'])

    return (' AND '.join(where) if where else '1=1'), params, {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'proveedor': proveedor,
    }



@admin_required
def admin_notas_credito_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _notas_credito_query, _notas_credito_resumen
    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_filtros(request)
    filas = _notas_credito_query(where, params.copy(), cliente)
    resumen = _notas_credito_resumen(where, params.copy(), cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = 'N_C Recibidas'
    headers = ['N°','PROVEEDOR','RUC','TIPO','FECHA','N° NOTA','AUTORIZACIÓN',
               'BASES SIN IVA','BASES CON IVA','IVA','TOTAL','COD MOD',
               'DOC. MODIFICADO','AUT. MODIFICACIÓN']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')
    for row in filas:
        ws.append(list(row))
    ws.append([])
    ws.append(['','','','','','','RESUMEN',
               resumen['bases_sin_iva'], resumen['bases_con_iva'],
               resumen['iva'], resumen['total']])
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="notas_credito_recibidas_administrativo.xlsx"'
    return response


@admin_required
def admin_notas_credito_pdf(request, ruc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .views import _notas_credito_query, _notas_credito_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_filtros(request)

    try:
        filas = _notas_credito_query(where, params.copy(), cliente)
        resumen = _notas_credito_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(
            f'Error generando PDF de notas de crédito recibidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    def _num(value):
        try:
            return f"{float(value or 0):.2f}"
        except (TypeError, ValueError):
            return "0.00"

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18,
    )

    styles = getSampleStyleSheet()
    titulo = ParagraphStyle(
        'AdminNCTitulo', parent=styles['Title'], fontName='Helvetica-Bold',
        fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3,
    )
    cab = ParagraphStyle(
        'AdminNCCab', parent=styles['Normal'], fontName='Helvetica-Bold',
        fontSize=8, leading=10, alignment=TA_CENTER,
    )
    tabla = ParagraphStyle(
        'AdminNCTabla', parent=styles['Normal'], fontName='Helvetica',
        fontSize=5, leading=5.8, alignment=TA_CENTER, wordWrap='CJK',
    )
    izq = ParagraphStyle('AdminNCIzq', parent=tabla, alignment=0)
    enc = ParagraphStyle(
        'AdminNCEnc', parent=tabla, fontName='Helvetica-Bold',
        textColor=colors.white, leading=6,
    )

    headers = [
        'N°', 'PROVEEDOR', 'RUC', 'TIPO', 'FECHA', 'N° NOTA',
        'AUTORIZACIÓN', 'BASES SIN IVA', 'BASES CON IVA', 'IVA',
        'TOTAL', 'COD MOD', 'DOC. MODIFICADO', 'AUT. MODIFICACIÓN',
    ]

    data = [[Paragraph(label, enc) for label in headers]]

    for row in filas:
        values = []
        for i, value in enumerate(row):
            if i in (1, 2, 3, 4, 5, 6, 11, 12, 13):
                values.append(
                    Paragraph(str(value or ''), izq if i == 1 else tabla)
                )
            else:
                values.append(Paragraph(_num(value), tabla))
        data.append(values)

    data.append([
        '', '', '', '', '', '', 'RESUMEN',
        _num(resumen['bases_sin_iva']),
        _num(resumen['bases_con_iva']),
        _num(resumen['iva']),
        _num(resumen['total']),
        '', '', '',
    ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[22, 105, 65, 38, 48, 70, 70, 55, 55, 45, 55, 45, 75, 75],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (7, 1), (10, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
        ('FONTNAME', (6, -1), (10, -1), 'Helvetica-Bold'),
    ]))

    elements = [
        Paragraph('REPORTE DE NOTAS DE CRÉDITO RECIBIDAS - ADMINISTRATIVO', titulo),
        Paragraph(
            f"DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}",
            cab,
        ),
        Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}", cab),
        Spacer(1, 8),
        table,
    ]

    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = (
        'attachment; filename="notas_credito_recibidas_administrativo.pdf"'
    )
    return response

@admin_required
def admin_notas_credito_recibidas(request, ruc):
    from .views import _cliente_db, _notas_credito_base_sql, _notas_credito_query, _notas_credito_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_filtros(request)
    db = _cliente_db(cliente)

    try:
        with db.cursor() as cursor:
            cursor.execute(
                f"SELECT COUNT(*) FROM ({_notas_credito_base_sql()}) notas_reporte WHERE {where}",
                params,
            )
            total_registros = cursor.fetchone()[0]
    except Exception as exc:
        return HttpResponse(
            f'Error consultando notas de crédito recibidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    try:
        pagina = max(1, int(request.GET.get('pagina', '1')))
    except (TypeError, ValueError):
        pagina = 1

    por_pagina = 50
    total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)
    pagina = min(pagina, total_paginas)
    offset = (pagina - 1) * por_pagina

    try:
        filas = _notas_credito_query(where, params.copy(), cliente, por_pagina, offset)
        resumen = _notas_credito_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(
            f'Error consultando notas de crédito recibidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    return render(request, 'admin/notas_credito_recibidas.html', {
        'cliente': cliente,
        'filas': filas,
        'resumen': resumen,
        'filtros': filtros,
        'pagina': pagina,
        'total_paginas': total_paginas,
        'total_registros': total_registros,
    })


def _admin_notas_credito_emitidas_filtros(request):
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)
    fecha_desde = request.GET.get('fecha_desde', '').strip() or primer_dia_mes.strftime('%Y-%m-%d')
    fecha_hasta = request.GET.get('fecha_hasta', '').strip() or hoy.strftime('%Y-%m-%d')

    where = []
    params = []

    try:
        datetime.strptime(fecha_desde, '%Y-%m-%d')
        where.append('fecnc::date >= %s::date')
        params.append(fecha_desde)
    except ValueError:
        fecha_desde = ''

    try:
        datetime.strptime(fecha_hasta, '%Y-%m-%d')
        where.append("fecnc::date < (%s::date + INTERVAL '1 day')")
        params.append(fecha_hasta)
    except ValueError:
        fecha_hasta = ''

    return (' AND '.join(where) if where else '1=1'), params, {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
    }


@admin_required
def admin_notas_credito_emitidas(request, ruc):
    from .views import _cliente_db
    from .ingresos_views import _ingresos_nc_base_sql, _ingresos_nc_query, _ingresos_nc_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_emitidas_filtros(request)
    db = _cliente_db(cliente)

    try:
        with db.cursor() as cursor:
            cursor.execute(
                f"SELECT COUNT(*) FROM ({_ingresos_nc_base_sql()}) nc_reporte WHERE {where}",
                params,
            )
            total_registros = cursor.fetchone()[0]
    except Exception as exc:
        return HttpResponse(
            f'Error consultando notas de crédito emitidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    try:
        pagina = max(1, int(request.GET.get('pagina', '1')))
    except (TypeError, ValueError):
        pagina = 1

    por_pagina = 50
    total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)
    pagina = min(pagina, total_paginas)
    offset = (pagina - 1) * por_pagina

    try:
        filas = _ingresos_nc_query(where, params.copy(), cliente, por_pagina, offset)
        resumen = _ingresos_nc_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(
            f'Error consultando notas de crédito emitidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    return render(request, 'admin/notas_credito_emitidas.html', {
        'cliente': cliente,
        'filas': filas,
        'resumen': resumen,
        'filtros': filtros,
        'pagina': pagina,
        'total_paginas': total_paginas,
        'total_registros': total_registros,
    })


@admin_required
def admin_notas_credito_emitidas_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .ingresos_views import _ingresos_nc_query, _ingresos_nc_resumen, INGRESOS_NC_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_emitidas_filtros(request)
    filas = _ingresos_nc_query(where, params.copy(), cliente)
    resumen = _ingresos_nc_resumen(where, params.copy(), cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = 'N_C Emitidas'
    ws.append([label for _, label in INGRESOS_NC_COLUMNS])

    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')

    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', 'RESUMEN',
               resumen['bases_sin_iva'], resumen['bases_con_iva'],
               resumen['iva'], resumen['total']])

    for column, width in enumerate([7, 35, 17, 13, 22, 28, 16, 16, 13, 16, 22, 18], 1):
        ws.column_dimensions[chr(64 + column)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="notas_credito_emitidas_administrativo.xlsx"'
    return response


@admin_required
def admin_notas_credito_emitidas_pdf(request, ruc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .ingresos_views import _ingresos_nc_query, _ingresos_nc_resumen, INGRESOS_NC_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_notas_credito_emitidas_filtros(request)

    try:
        filas = _ingresos_nc_query(where, params.copy(), cliente)
        resumen = _ingresos_nc_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(
            f'Error generando PDF de notas de crédito emitidas: {type(exc).__name__}: {exc}',
            status=500,
            content_type='text/plain; charset=utf-8',
        )

    def _num(value):
        try:
            return f"{float(value or 0):.2f}"
        except (TypeError, ValueError):
            return "0.00"

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18,
    )
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle(
        'AdminNCEEmitTitulo', parent=styles['Title'], fontName='Helvetica-Bold',
        fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3,
    )
    cab = ParagraphStyle(
        'AdminNCEEmitCab', parent=styles['Normal'], fontName='Helvetica-Bold',
        fontSize=8, leading=10, alignment=TA_CENTER,
    )
    tabla = ParagraphStyle(
        'AdminNCEEmitTabla', parent=styles['Normal'], fontName='Helvetica',
        fontSize=5.2, leading=6, alignment=TA_CENTER, wordWrap='CJK',
    )
    izq = ParagraphStyle('AdminNCEEmitIzq', parent=tabla, alignment=0)
    enc = ParagraphStyle(
        'AdminNCEEmitEnc', parent=tabla, fontName='Helvetica-Bold',
        textColor=colors.white, leading=6,
    )

    data = [[Paragraph(label, enc) for _, label in INGRESOS_NC_COLUMNS]]
    for row in filas:
        values = []
        for i, value in enumerate(row):
            if i in (1, 2, 3, 4, 5, 10, 11):
                values.append(Paragraph(str(value or ''), izq if i == 1 else tabla))
            else:
                values.append(Paragraph(_num(value), tabla))
        data.append(values)

    data.append([
        '', '', '', '', '', '', _num(resumen['bases_sin_iva']),
        _num(resumen['bases_con_iva']), _num(resumen['iva']),
        _num(resumen['total']), '', '',
    ])

    table = Table(
        data, repeatRows=1,
        colWidths=[22, 100, 68, 48, 70, 78, 58, 58, 45, 55, 65, 60],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (6, 1), (9, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
        ('FONTNAME', (6, -1), (9, -1), 'Helvetica-Bold'),
    ]))

    doc.build([
        Paragraph('REPORTE DE NOTAS DE CRÉDITO EMITIDAS - ADMINISTRATIVO', titulo),
        Paragraph(
            f"DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}",
            cab,
        ),
        Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}", cab),
        Spacer(1, 8),
        table,
    ])

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="notas_credito_emitidas_administrativo.pdf"'
    return response


@admin_required
def admin_compras(request, ruc):
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    db = _cliente_db(cliente)

    try:
        with db.cursor() as cursor:
            cursor.execute(
                f"SELECT COUNT(*) FROM ({_compras_base_sql()}) compras_reporte WHERE {where}",
                params,
            )
            total_registros = cursor.fetchone()[0]
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    try:
        pagina = max(1, int(request.GET.get('pagina', '1')))
    except (TypeError, ValueError):
        pagina = 1

    por_pagina = 50
    total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)
    pagina = min(pagina, total_paginas)
    offset = (pagina - 1) * por_pagina

    try:
        filas = _compras_query(where, params.copy(), cliente, por_pagina, offset, include_numcompra=True)
        resumen = _compras_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    return render(request, 'admin/compras.html', {
        'cliente': cliente,
        'filas': filas,
        'resumen': resumen,
        'filtros': filtros,
        'pagina': pagina,
        'total_paginas': total_paginas,
        'total_registros': total_registros,
    })


@admin_required
def admin_compras_editar(request, ruc):
    """Carga y actualiza una compra desde el módulo administrativo."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    numcompra = request.POST.get('numcompra', '').strip() if request.method == 'POST' else request.GET.get('numcompra', '').strip()
    autorizacion = request.POST.get('autorizacion', '').strip() if request.method == 'POST' else request.GET.get('autorizacion', '').strip()
    proveedor_ruc = request.POST.get('proveedor_ruc', '').strip() if request.method == 'POST' else request.GET.get('proveedor_ruc', '').strip()
    documento = request.POST.get('documento', '').strip() if request.method == 'POST' else request.GET.get('documento', '').strip()

    try:
        # Conexión local al PostgreSQL de la base del cliente.
        # Se construye aquí para que el módulo administrativo no dependa
        # de la función privada de views.py.
        alias = f'cliente_{cliente.ruccedcli}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = cliente.ruccedcli
            connections.databases[alias] = base
        db = connections[alias]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error conectando a la base del cliente: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor_ctx = db.cursor()
        cursor = cursor_ctx.__enter__()
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error abriendo cursor de compras: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor.execute("""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema='public' AND table_name='comprasnue'
            ORDER BY ordinal_position
        """)
        columnas = [r[0] for r in cursor.fetchall()]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error consultando estructura de comprasnue: {type(exc).__name__}: {exc}'
        }, status=500)

    if not columnas:
        return JsonResponse({'ok': False, 'error': 'No se encontró la tabla comprasnue.'}, status=404)
    
    cols = set(columnas)
    # Identifica, sin asumir un nombre único, los campos usados por instalaciones antiguas.
    mes_col = next((x for x in ('mesdeclaracion','mesdeclara','mesdec','mes_declaracion','mes') if x in cols), None)
    anio_col = next((x for x in ('aniodeclaracion','aniodeclara','aniodec','anio_declaracion','anio','ano','año') if x in cols), None)
    
    select_cols = [
        'numcompra','fecemi','ruccedprovee','nomprovee','tipcom','numest','numptoemi','numsec','numaut',
        'baseimpnoobj','baseimpiva0','baseexenta','baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    select_cols = [x for x in select_cols if x in cols]
    if mes_col: select_cols.append(mes_col)
    if anio_col: select_cols.append(anio_col)
    
    if 'numcompra' not in cols:
        return JsonResponse({'ok': False, 'error': 'El campo numcompra no existe en comprasnue.'}, status=500)
    if numcompra:
        where = '"numcompra"::text=%s'
        params = [numcompra]
    elif proveedor_ruc and documento:
        partes = documento.split('-')
        if len(partes) != 3:
            return JsonResponse({'ok': False, 'error': 'Número de factura inválido.'}, status=400)
        numest, numptoemi, numsec = [p.strip() for p in partes]
        where = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
        params = [proveedor_ruc, numest, numptoemi, numsec]
    else:
        return JsonResponse({'ok': False, 'error': 'No se recibió el identificador de la compra.'}, status=400)
    
    try:
        select_sql = f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where} LIMIT 1'
        cursor.execute(select_sql, params)
        row = cursor.fetchone()
    
        if not row and numcompra and proveedor_ruc and documento:
            partes = documento.split('-')
            if len(partes) == 3:
                numest, numptoemi, numsec = [p.strip() for p in partes]
                where_fallback = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
                params_fallback = [proveedor_ruc, numest, numptoemi, numsec]
                cursor.execute(
                    f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where_fallback} LIMIT 1',
                    params_fallback
                )
                row = cursor.fetchone()
    except Exception as exc:
        db.rollback()
        return JsonResponse({'ok': False, 'error': f'Error consultando la compra: {type(exc).__name__}: {exc}'}, status=500)
    if not row:
        return JsonResponse({
            'ok': False,
            'error': f'No se encontró la compra. numcompra recibido: {numcompra or "(vacío)"}; RUC: {proveedor_ruc or "(vacío)"}; factura: {documento or "(vacía)"}'
        }, status=404)
    
    data = dict(zip(select_cols, row))
    if request.method == 'GET':
        for k, v in list(data.items()):
            if hasattr(v, 'isoformat'):
                data[k] = v.isoformat()
            elif v is None:
                data[k] = ''
            else:
                data[k] = str(v)
        data['_mes_col'] = mes_col or ''
        data['_anio_col'] = anio_col or ''
        return JsonResponse({'ok': True, 'compra': data})
    
    editable = [
        'fecemi','tipcom','numaut','baseimpnoobj','baseimpiva0','baseexenta',
        'baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    values = {}
    for field in editable:
        if field in cols and field in request.POST:
            values[field] = request.POST.get(field, '').strip() or None
    # IVA se recalcula en servidor a partir de los subtotales.
    tasas = {'baseimpiva5': ('montoiva5', 0.05), 'baseimpiva8': ('montoiva8', 0.08),
             'baseimpiva12': ('montoiva12', 0.12), 'baseimpiva14': ('montoiva14', 0.14),
             'baseimpiva15': ('montoiva15', 0.15)}
    for base, (iva, tasa) in tasas.items():
        if base in cols:
            try:
                base_val = float(values.get(base) or 0)
                values[iva] = round(base_val * tasa, 2)
            except (TypeError, ValueError):
                values[iva] = 0
    
    if mes_col and 'mes_declaracion' in request.POST:
        values[mes_col] = request.POST.get('mes_declaracion') or None
    if anio_col and 'anio_declaracion' in request.POST:
        values[anio_col] = request.POST.get('anio_declaracion') or None
    
    if not values:
        return JsonResponse({'ok': False, 'error': 'No hay cambios para guardar.'}, status=400)
    
    sets = ', '.join(f'"{k}"=%s' for k in values)
    cursor.execute(f'UPDATE comprasnue SET {sets} WHERE {where}', [*values.values(), *params])
    db.commit()
    return JsonResponse({'ok': True, 'message': 'Compra actualizada correctamente.'})


@admin_required
def admin_compras_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compras'
    ws.append([label for _, label in COMPRAS_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')
    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', '', 'RESUMEN',
               resumen['bases_sin_iva'], resumen['bases_con_iva'], resumen['iva'],
               resumen['total'], resumen['retiva'], '', resumen['retrenta'], ''])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    widths = [7, 38, 17, 10, 13, 25, 18, 16, 16, 13, 16, 13, 11, 16, 18]
    for i, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.xlsx"'
    return response


@admin_required
def admin_compras_pdf(request, ruc):
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .views import _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('AdminComprasPDFTitulo', parent=styles['Title'], fontName='Helvetica-Bold',
                            fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3)
    cab = ParagraphStyle('AdminComprasPDFCab', parent=styles['Normal'], fontName='Helvetica-Bold',
                         fontSize=8, leading=10, alignment=TA_CENTER)
    tabla = ParagraphStyle('AdminComprasPDFTabla', parent=styles['Normal'], fontName='Helvetica',
                           fontSize=5, leading=5.8, alignment=TA_CENTER, wordWrap='CJK')
    izq = ParagraphStyle('AdminComprasPDFIzq', parent=tabla, alignment=0)
    enc = ParagraphStyle('AdminComprasPDFEnc', parent=tabla, fontName='Helvetica-Bold',
                         textColor=colors.white, leading=6)

    elements = [
        Paragraph('REPORTE DE COMPRAS - ADMINISTRATIVO', titulo),
        Paragraph(f"DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}", cab),
        Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}", cab),
        Spacer(1, 8),
    ]

    data = [[Paragraph(label, enc) for _, label in COMPRAS_COLUMNS]]
    for row in filas:
        values = []
        for i, value in enumerate(row):
            if i in (1, 2, 3, 4, 5, 6, 12, 14):
                values.append(Paragraph(str(value or ''), izq if i == 1 else tabla))
            else:
                values.append(Paragraph(f"{float(value or 0):.2f}", tabla))
        data.append(values)

    data.append(['', '', '', '', '', '', '', f"{resumen['bases_sin_iva']:.2f}",
                 f"{resumen['bases_con_iva']:.2f}", f"{resumen['iva']:.2f}",
                 f"{resumen['total']:.2f}", f"{resumen['retiva']:.2f}", '',
                 f"{resumen['retrenta']:.2f}", ''])

    table = Table(data, repeatRows=1,
                  colWidths=[22, 105, 65, 38, 48, 70, 70, 50, 50, 40, 50, 42, 38, 48, 52])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (7, 1), (11, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.pdf"'
    return response


@admin_required
def admin_clientes(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    clientes = Cliente.objects.all().order_by('nomclient')

    if query:
        from django.db.models import Q
        clientes = clientes.filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if estado == 'activos':
        clientes = clientes.filter(activo=True)
    elif estado == 'inactivos':
        clientes = clientes.filter(activo=False)

    if dia != 'todos':
        clientes = clientes.filter(diadeclaracion=dia)

    if tipdec:
        clientes = clientes.filter(tipdec__iexact=tipdec)

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/clientes.html',
        {
            'clientes': clientes,
            'query': query,
            'estado': estado,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
        },
    )


@admin_required
def admin_suscriptores(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()

    suscriptores = Suscriptor.objects.all()

    if query:
        from django.db.models import Q
        suscriptores = suscriptores.filter(
            Q(nombre__icontains=query) |
            Q(email__icontains=query)
        )

    if estado == 'activos':
        suscriptores = suscriptores.filter(activo=True)
    elif estado == 'inactivos':
        suscriptores = suscriptores.filter(activo=False)

    return render(
        request,
        'admin/suscriptores.html',
        {
            'suscriptores': suscriptores,
            'query': query,
            'estado': estado,
            'suscriptores_total': Suscriptor.objects.count(),
            'suscriptores_activos': Suscriptor.objects.filter(activo=True).count(),
        },
    )


@admin_required
def admin_contrasenas(request):
    """Listado administrativo de clientes activos y sus credenciales tributarias/laborales."""
    query = request.GET.get('q', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    # Al entrar se muestran únicamente clientes activos.
    # Cuando se realiza una búsqueda, se consulta toda la base de datos,
    # incluyendo clientes inactivos.
    clientes_qs = Cliente.objects.filter(activo=True)

    if query:
        from django.db.models import Q
        clientes_qs = Cliente.objects.all().filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if dia != 'todos':
        clientes_qs = clientes_qs.filter(diadeclaracion=str(dia))

    if tipdec:
        clientes_qs = clientes_qs.filter(semensual__iexact=tipdec)

    clientes_qs = clientes_qs.order_by('nomclient')

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/contrasenas.html',
        {
            'clientes': clientes_qs,
            'pestanas': [
                {
                    'dia': dia_item,
                    'clientes': clientes_qs.filter(diadeclaracion=str(dia_item)),
                }
                for dia_item in dias
            ],
            'clientes_busqueda': clientes_qs if query else Cliente.objects.none(),
            'query': query,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


@admin_required
def admin_cliente(request, ruc):
    cliente = get_object_or_404(Cliente, pk=ruc)

    if request.method == 'POST' and request.POST.get('accion') == 'guardar_datos_cliente':
        cliente.dirclient = request.POST.get('dirclient', '').strip()
        cliente.teldomcli = request.POST.get('teldomcli', '').strip()
        cliente.teloficli = request.POST.get('teloficli', '').strip()
        cliente.telcelcli = request.POST.get('telcelcli', '').strip()
        cliente.corelectr = request.POST.get('corelectr', '').strip()
        cliente.ocuclient = request.POST.get('ocuclient', '').strip()
        cliente.save(update_fields=[
            'dirclient', 'teldomcli', 'teloficli',
            'telcelcli', 'corelectr', 'ocuclient',
        ])

        messages.success(request, 'Datos básicos del cliente actualizados correctamente.')

        return redirect('admin_cliente', ruc=cliente.ruccedcli)

    usuarios = (
        UsuarioCliente.objects
        .filter(cliente=cliente)
        .select_related('usuario')
        .order_by('usuario__username')
    )

    db_name = str(cliente.ruccedcli).strip()
    from .models.archivo import Archivo
    from .services.archivos.storage import data_root
    certificado_ruc_path = data_root() / str(cliente.ruccedcli).strip() / 'documentos' / 'ruc' / 'certificado_ruc.pdf'
    certificado_ruc_exists = (
        certificado_ruc_path.is_file()
        and Archivo.objects.filter(
            cliente=cliente,
            tipo=Archivo.Tipo.PDF,
            ruta_relativa=f'{str(cliente.ruccedcli).strip()}/documentos/ruc/certificado_ruc.pdf',
            activo=True,
        ).exists()
    )
    db_status = 'No verificada'
    db_error = ''
    conciliacion = {
        'anio': 0,
        'anios': [],
        'meses': [],
        'totales': {
            'ventas': 0,
            'compras': 0,
            'retrenta': 0,
            'retiva': 0,
            'resultado': 0,
            'retenciones': 0,
        },
    }

    # La conexión y la consulta de conciliación se manejan por separado.
    # Un error en una consulta no debe hacer aparecer la base como desconectada.
    try:
        alias = f'cliente_{db_name}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = db_name
            connections.databases[alias] = base
        client_connection = connections[alias]
        client_connection.ensure_connection()
        db_status = 'Conectada'
    except Exception as exc:
        db_status = 'No disponible'
        db_error = str(exc)

    if db_status == 'Conectada':
        try:
            conciliacion = _admin_conciliacion_cliente(cliente, request.GET.get('anio'))
        except Exception as exc:
            db_error = f'Error consultando datos: {exc}'

    return render(
        request,
        'admin/cliente.html',
        {
            'cliente': cliente,
            'usuarios': usuarios,
            'db_name': db_name,
            'db_status': db_status,
            'db_error': db_error,
            'conciliacion': conciliacion,
            'certificado_ruc_exists': certificado_ruc_exists,
        },
    )


@admin_required
def admin_cliente_certificado_ruc(request, ruc):
    cliente = get_object_or_404(Cliente, pk=ruc)
    if request.method == 'POST':
        certificado_ruc = request.FILES.get('certificado_ruc')
        if not certificado_ruc:
            messages.error(request, 'Selecciona un archivo PDF de Certificado de RUC.')
            return redirect('admin_cliente', ruc=cliente.ruccedcli)

        try:
            if certificado_ruc.size > 10 * 1024 * 1024:
                raise ValueError('El Certificado de RUC no puede superar los 10 MB.')

            extension = Path(certificado_ruc.name).suffix.lower()
            content_type = (certificado_ruc.content_type or '').lower()
            if extension != '.pdf' or content_type not in ('application/pdf', 'application/octet-stream'):
                raise ValueError('El Certificado de RUC debe ser un archivo PDF.')

            if certificado_ruc.size <= 0:
                raise ValueError('El archivo PDF está vacío.')

            from .models.archivo import Archivo
            from .services.archivos.storage import guardar_archivo, eliminar_archivo

            anteriores = Archivo.objects.filter(
                cliente=cliente,
                tipo=Archivo.Tipo.PDF,
                ruta_relativa__startswith=f'{cliente.ruccedcli}/documentos/ruc/'
            )
            for anterior in anteriores:
                eliminar_archivo(anterior)

            guardar_archivo(
                cliente=cliente,
                uploaded_file=certificado_ruc,
                tipo=Archivo.Tipo.PDF,
                subcarpeta='documentos/ruc',
                nombre_fisico='certificado_ruc.pdf',
                usuario=request.user if getattr(request.user, 'is_authenticated', False) else None,
            )
            messages.success(request, 'Certificado de RUC cargado correctamente.')
        except Exception as exc:
            messages.error(request, f'No se pudo guardar el Certificado de RUC: {type(exc).__name__}: {exc}')

        return redirect('admin_cliente', ruc=cliente.ruccedcli)

    archivo = Path(settings.TOTALCOUNTS_DATA_ROOT) / str(cliente.ruccedcli) / 'documentos' / 'ruc' / 'certificado_ruc.pdf'
    if not archivo.is_file():
        raise Http404('No existe un Certificado de RUC cargado para este cliente.')
    response = FileResponse(archivo.open('rb'), content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="certificado_ruc.pdf"'
    return response


@admin_required
def admin_cliente_documento_legal(request, ruc, tipo):
    cliente = get_object_or_404(Cliente, pk=ruc)

    documentos = {
        'contrato_constitutivo': {
            'subcarpeta': 'documentos/legales',
            'nombre': 'contrato_constitutivo.pdf',
            'titulo': 'Contrato Constitutivo',
        },
        'nombramiento_rep_legal': {
            'subcarpeta': 'documentos/legales',
            'nombre': 'nombramiento_rep_legal.pdf',
            'titulo': 'Nombramiento Rep. Legal',
        },
        'nomina_socios': {
            'subcarpeta': 'documentos/legales',
            'nombre': 'nomina_socios.pdf',
            'titulo': 'Nómina de Socios',
        },
    }
    documento = documentos.get(tipo)
    if not documento:
        raise Http404('Documento legal no válido.')

    from .models.archivo import Archivo
    from .services.archivos.storage import guardar_archivo, eliminar_archivo

    ruc_cliente = str(cliente.ruccedcli).strip()
    ruta_relativa = f'{ruc_cliente}/{documento["subcarpeta"]}/{documento["nombre"]}'

    if request.method == 'POST':
        archivo_subido = request.FILES.get('documento_legal')
        if not archivo_subido:
            messages.error(request, f'Selecciona el archivo de {documento["titulo"]}.')
            return redirect('admin_cliente', ruc=cliente.ruccedcli)

        try:
            if archivo_subido.size <= 0:
                raise ValueError('El archivo está vacío.')
            if archivo_subido.size > 10 * 1024 * 1024:
                raise ValueError('El archivo no puede superar los 10 MB.')

            extension = Path(archivo_subido.name).suffix.lower()
            content_type = (archivo_subido.content_type or '').lower()
            if extension != '.pdf' or content_type not in ('application/pdf', 'application/octet-stream'):
                raise ValueError(f'{documento["titulo"]} debe ser un archivo PDF.')

            anteriores = Archivo.objects.filter(
                cliente=cliente,
                tipo=Archivo.Tipo.PDF,
                ruta_relativa=ruta_relativa,
            )
            for anterior in anteriores:
                eliminar_archivo(anterior)

            guardar_archivo(
                cliente=cliente,
                uploaded_file=archivo_subido,
                tipo=Archivo.Tipo.PDF,
                subcarpeta=documento['subcarpeta'],
                nombre_fisico=documento['nombre'],
                usuario=request.user if getattr(request.user, 'is_authenticated', False) else None,
            )
            messages.success(request, f'{documento["titulo"]} cargado correctamente.')
        except Exception as exc:
            messages.error(request, f'No se pudo guardar {documento["titulo"]}: {type(exc).__name__}: {exc}')

        return redirect('admin_cliente', ruc=cliente.ruccedcli)

    ruta = Path(settings.TOTALCOUNTS_DATA_ROOT) / ruc_cliente / documento['subcarpeta'] / documento['nombre']
    if not ruta.is_file():
        raise Http404(f'No existe {documento["titulo"]}.')

    response = FileResponse(ruta.open('rb'), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{documento["nombre"]}"'
    return response



def _conta_sincronizar_compras(usuario, ruc, anio, mes, tipo_comprobante=1):
    """Llama a Conta desde el servidor; el navegador nunca recibe el secreto."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
        'tipo_comprobante': int(tipo_comprobante),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc




def _conta_sincronizar_ventas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de facturas emitidas hacia ventas."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/ventas/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_validar_ventas(usuario, ruc, anio, mes):
    """Inicia en Conta la validación día por día de ventas contra el SRI."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/ventas/validar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_sincronizar_notas_credito_emitidas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de notas de crédito emitidas."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc,
        'anio': int(anio),
        'mes': int(mes),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/ventas/notas-credito/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_sincronizar_retenciones_emitidas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de retenciones emitidas hacia comprasnue."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'ruc': ruc, 'anio': int(anio), 'mes': int(mes)}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/retenciones-emitidas/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=180) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_sincronizar_notas_credito_recibidas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de notas de crédito recibidas hacia comprasnue."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'ruc': ruc, 'anio': int(anio), 'mes': int(mes)}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/notas-credito/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


def _conta_sri_status(usuario, job_id):
    """Consulta el estado de un trabajo SRI ya iniciado en Conta."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/estado/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='GET',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


def _conta_sri_cancelar(usuario, job_id):
    """Solicita a Conta detener un trabajo SRI."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/cancelar/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_chat(usuario, mensaje, conversation_id='admin'):
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'mensaje': mensaje, 'conversation_id': conversation_id or 'admin'}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/chat',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=180) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


def _conta_sincronizar_retenciones_recibidas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de retenciones recibidas hacia ventas."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'ruc': ruc, 'anio': int(anio), 'mes': int(mes)}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/retenciones/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=180) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


@admin_required
def admin_conta_sri_cancelar(request):
    """Solicita la cancelación de un trabajo SRI en Conta."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.POST.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_cancelar(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta_sri_status(request):
    """Endpoint AJAX del panel para consultar trabajos SRI en segundo plano."""
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.GET.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_status(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta(request):
    """Panel de Conta dentro del administrador autenticado de TotalCounts."""
    from datetime import datetime

    clientes = Cliente.objects.filter(activo=True).order_by('nomclient')
    ruc = request.POST.get('ruc', '').strip() if request.method == 'POST' else request.GET.get('ruc', '').strip()
    anio = request.POST.get('anio', '').strip() if request.method == 'POST' else request.GET.get('anio', '').strip()
    mes = request.POST.get('mes', '').strip() if request.method == 'POST' else request.GET.get('mes', '').strip()

    if not anio:
        anio = str(datetime.now().year)
    if not mes:
        mes = str(datetime.now().month)

    resultado = None
    sri_job = None
    respuesta_ia = None
    error = ''

    if request.method == 'POST' and request.POST.get('accion') == 'chat':
        mensaje = request.POST.get('mensaje', '').strip()
        if not mensaje:
            error = 'Escriba una consulta para Conta.'
        else:
            try:
                respuesta_ia = _conta_chat(
                    usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
                    mensaje=mensaje,
                    conversation_id=request.session.session_key or 'admin',
                )
            except RuntimeError as exc:
                error = str(exc)

    elif request.method == 'POST' and request.POST.get('accion') in ('sincronizar_compras', 'sincronizar_ventas', 'sincronizar_notas_credito_emitidas', 'sincronizar_notas_credito_recibidas', 'sincronizar_retenciones_recibidas', 'sincronizar_retenciones_emitidas', 'validar_ventas'):
        cliente = clientes.filter(ruccedcli=ruc).first()
        if cliente is None:
            error = 'Seleccione un cliente activo válido.'
        else:
            try:
                usuario = request.session.get(ADMIN_USERNAME_KEY, 'ADMIN')
                if request.POST.get('accion') == 'sincronizar_ventas':
                    resultado = _conta_sincronizar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de ventas emitidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'validar_ventas':
                    resultado = _conta_validar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La validación de ventas contra el SRI fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_notas_credito_recibidas':
                    resultado = _conta_sincronizar_notas_credito_recibidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de notas de crédito recibidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_retenciones_recibidas':
                    resultado = _conta_sincronizar_retenciones_recibidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de retenciones recibidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_retenciones_emitidas':
                    resultado = _conta_sincronizar_retenciones_emitidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de retenciones emitidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_notas_credito_emitidas':
                    resultado = _conta_sincronizar_notas_credito_emitidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de notas de crédito emitidas fue iniciada en segundo plano.')
                else:
                    resultado = _conta_sincronizar_compras(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                        tipo_comprobante=int(request.POST.get('tipo_comprobante', '1')),
                    )
                    messages.success(request, 'La sincronización de compras fue iniciada en segundo plano.')
                sri_job = resultado.get('job_id')
            except (ValueError, RuntimeError) as exc:
                error = str(exc)

    return render(
        request,
        'admin/conta.html',
        {
            'clientes': clientes,
            'ruc': ruc,
            'anio': anio,
            'mes': mes,
            'resultado': resultado,
            'sri_job': sri_job,
            'respuesta_ia': respuesta_ia,
            'error': error,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )

def _admin_conciliacion_cliente(cliente, anio=None):
    """Obtiene la conciliación mensual de ventas y compras."""
    db_name = str(cliente.ruccedcli).strip()
    alias = f'cliente_{db_name}'
    if alias not in connections.databases:
        base = settings.DATABASES['default'].copy()
        base['NAME'] = db_name
        connections.databases[alias] = base

    db = connections[alias]

    with db.cursor() as cursor:
        cursor.execute("""
            SELECT DISTINCT anio
            FROM (
                SELECT EXTRACT(YEAR FROM fecfactur::date)::integer AS anio
                FROM ventas
                WHERE fecfactur IS NOT NULL
                UNION
                SELECT TRIM("año"::text)::integer AS anio
                FROM comprasnue
                WHERE TRIM("año"::text) ~ '^\\d{4}$'
            ) periodos
            WHERE anio IS NOT NULL
            ORDER BY anio DESC
        """)
        anios = [int(row[0]) for row in cursor.fetchall()]

    if not anios:
        return {
            'anio': anio or 0, 'anios': [], 'meses': [],
            'mes_final_nombre': '',
            'totales': {
                'ventas': 0, 'compras': 0, 'retrenta': 0,
                'retiva': 0, 'resultado': 0, 'retenciones': 0,
            },
        }

    try:
        anio = int(anio)
    except (TypeError, ValueError):
        anio = anios[0]

    if anio not in anios:
        anio = anios[0]

    with db.cursor() as cursor:
        cursor.execute("""
            WITH meses AS (
                SELECT generate_series(1, 12) AS mes
            ),
            ventas_mes AS (
                SELECT
                    EXTRACT(MONTH FROM fecfactur::date)::integer AS mes,
                    COALESCE(SUM(
                        COALESCE(NULLIF(basenoobj::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva0::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva5::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva8::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva12::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva14::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseiva15::text, ''), '0')::numeric
                    ), 0) AS ventas,
                    COALESCE(SUM(COALESCE(NULLIF(retrenta::text, ''), '0')::numeric), 0) AS retrenta,
                    COALESCE(SUM(COALESCE(NULLIF(retiva::text, ''), '0')::numeric), 0) AS retiva
                FROM ventas
                WHERE fecfactur IS NOT NULL
                  AND EXTRACT(YEAR FROM fecfactur::date)::integer = %s
                GROUP BY EXTRACT(MONTH FROM fecfactur::date)::integer
            ),
            compras_mes AS (
                SELECT
                    TRIM(mes::text)::integer AS mes,
                    COALESCE(SUM(
                        COALESCE(NULLIF(baseimpnoobj::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva0::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseexenta::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva5::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva8::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva12::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva14::text, ''), '0')::numeric +
                        COALESCE(NULLIF(baseimpiva15::text, ''), '0')::numeric
                    ), 0) AS compras
                FROM comprasnue
                WHERE TRIM("año"::text) = %s
                  AND TRIM(mes::text) ~ '^\\d{1,2}$'
                  AND TRIM(mes::text)::integer BETWEEN 1 AND 12
                  AND TRIM(tipcom::text) IN ('01', '02')
                GROUP BY TRIM(mes::text)::integer
            )
            SELECT
                m.mes,
                COALESCE(v.ventas, 0),
                COALESCE(c.compras, 0),
                COALESCE(v.retrenta, 0),
                COALESCE(v.retiva, 0)
            FROM meses m
            LEFT JOIN ventas_mes v ON v.mes = m.mes
            LEFT JOIN compras_mes c ON c.mes = m.mes
            ORDER BY m.mes
        """, [str(anio), str(anio)])
        rows = cursor.fetchall()

    nombres = [
        'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO',
        'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE'
    ]

    meses = []
    for mes, ventas, compras, retrenta, retiva in rows:
        meses.append({
            'nombre': nombres[int(mes) - 1],
            'ventas': ventas or 0,
            'compras': compras or 0,
            'retrenta': retrenta or 0,
            'retiva': retiva or 0,
        })

    totales = {
        campo: sum((m[campo] for m in meses), 0)
        for campo in ('ventas', 'compras', 'retrenta', 'retiva')
    }
    totales['resultado'] = totales['ventas'] - totales['compras']
    totales['retenciones'] = totales['retrenta'] + totales['retiva']

    return {
        'anio': anio,
        'anios': anios,
        'meses': meses,
        'mes_final_nombre': 'DICIEMBRE',
        'totales': totales,
    }


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


def _admin_compras_filtros(request):
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)
    fecha_desde = request.GET.get('fecha_desde', '').strip() or primer_dia_mes.strftime('%Y-%m-%d')
    fecha_hasta = request.GET.get('fecha_hasta', '').strip() or hoy.strftime('%Y-%m-%d')
    proveedor = request.GET.get('proveedor', '').strip()

    where = []
    params = []
    try:
        datetime.strptime(fecha_desde, '%Y-%m-%d')
        where.append('fecemi::date >= %s::date')
        params.append(fecha_desde)
    except ValueError:
        fecha_desde = ''

    try:
        datetime.strptime(fecha_hasta, '%Y-%m-%d')
        where.append("fecemi::date < (%s::date + INTERVAL '1 day')")
        params.append(fecha_hasta)
    except ValueError:
        fecha_hasta = ''

    if proveedor:
        where.append('(ruccedprovee ILIKE %s OR nomprovee ILIKE %s)')
        params.extend([f'%{proveedor}%', f'%{proveedor}%'])

    return (' AND '.join(where) if where else '1=1'), params, {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'proveedor': proveedor,
    }


@admin_required
def admin_ventas(request, ruc):
    from .views import _cliente_db, _ventas_base_sql, _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_ventas_filtros(request)
    try:
        base = _ventas_base_sql()
        with _cliente_db(cliente).cursor() as cursor:
            cursor.execute(f"SELECT COUNT(*) FROM ({base}) ventas_reporte WHERE {where}", params)
            total_registros = cursor.fetchone()[0]
        try:
            pagina = max(1, int(request.GET.get('pagina','1')))
        except ValueError:
            pagina = 1
        por_pagina = 50
        filas = _ventas_query(where, params.copy(), cliente, por_pagina, (pagina-1)*por_pagina)
        resumen = _ventas_resumen(where, params.copy(), cliente)
        total_paginas = max(1, (total_registros + por_pagina - 1)//por_pagina)
    except Exception as exc:
        return HttpResponse(f"Error en reporte de ventas: {type(exc).__name__}: {exc}", status=500, content_type='text/plain; charset=utf-8')
    return render(request, 'admin/ventas.html', {
        'cliente': cliente, 'filas': filas, 'resumen': resumen, 'filtros': filtros,
        'pagina': pagina, 'total_paginas': total_paginas, 'total_registros': total_registros,
        'ventas_columns': VENTAS_COLUMNS,
    })


def _admin_ventas_filtros(request):
    hoy = timezone.localdate()
    primer_dia = hoy.replace(day=1)
    desde = request.GET.get('fecha_desde','').strip() or primer_dia.strftime('%Y-%m-%d')
    hasta = request.GET.get('fecha_hasta','').strip() or hoy.strftime('%Y-%m-%d')
    cliente_busqueda = request.GET.get('cliente_busqueda','').strip()
    where=[]; params=[]
    try:
        datetime.strptime(desde,'%Y-%m-%d'); where.append('fecfactur::date >= %s::date'); params.append(desde)
    except ValueError: desde=''
    try:
        datetime.strptime(hasta,'%Y-%m-%d'); where.append("fecfactur::date < (%s::date + INTERVAL '1 day')"); params.append(hasta)
    except ValueError: hasta=''
    if cliente_busqueda:
        where.append('(ruccedcli ILIKE %s OR nomcli ILIKE %s)')
        params.extend([f'%{cliente_busqueda}%',f'%{cliente_busqueda}%'])
    return (' AND '.join(where) if where else '1=1'), params, {'fecha_desde':desde,'fecha_hasta':hasta,'cliente_busqueda':cliente_busqueda}


@admin_required
def admin_ventas_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _cliente_db, _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente=get_object_or_404(Cliente,pk=ruc); where,params,filtros=_admin_ventas_filtros(request)
    filas=_ventas_query(where,params.copy(),cliente); resumen=_ventas_resumen(where,params.copy(),cliente)
    wb=Workbook(); ws=wb.active; ws.title='Ventas'
    ws.append([x[1] for x in VENTAS_COLUMNS])
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='21333E'); cell.alignment=Alignment(horizontal='center')
    for row in filas: ws.append(list(row))
    ws.append([]); ws.append(['','','','','','RESUMEN',resumen['base0'],resumen['baseiva'],resumen['iva'],resumen['total'],resumen['retiva'],resumen['retrenta']])
    buffer=BytesIO(); wb.save(buffer)
    response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']='attachment; filename="ventas_administrativo.xlsx"'; return response


@admin_required
def admin_ventas_editar(request, ruc):
    from .views import _cliente_db
    cliente=get_object_or_404(Cliente,pk=ruc)
    numfactur=(request.POST.get('numfactur') if request.method=='POST' else request.GET.get('numfactur','')).strip()
    db=_cliente_db(cliente)
    try:
        with db.cursor() as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_schema='public' AND table_name='ventas' ORDER BY ordinal_position")
            cols=[x[0] for x in cur.fetchall()]
            if 'numfactur' not in cols: return JsonResponse({'ok':False,'error':'El campo numfactur no existe en ventas.'},status=500)
            wanted=['fecfactur','numfactur','autorizacion','basenoobj','baseiva0','baseiva12','iva','retiva','retrenta','numret','autret','codret']
            select=[x for x in wanted if x in cols]
            cur.execute(f'SELECT {",".join(chr(34)+x+chr(34) for x in select)} FROM ventas WHERE "numfactur"::text=%s LIMIT 1',[numfactur])
            row=cur.fetchone()
            if not row: return JsonResponse({'ok':False,'error':f'No se encontró la factura {numfactur}.'},status=404)
            data=dict(zip(select,row))
            if request.method=='GET':
                for k,v in data.items(): data[k]='' if v is None else (v.isoformat() if hasattr(v,'isoformat') else str(v))
                return JsonResponse({'ok':True,'venta':data})
            editable=[x for x in ['fecfactur','autorizacion','basenoobj','baseiva0','baseiva12','iva','retiva','retrenta','numret','autret','codret'] if x in cols and x in request.POST]
            vals={x:(request.POST.get(x,'').strip() or None) for x in editable}
            if not vals: return JsonResponse({'ok':False,'error':'No hay cambios para guardar.'},status=400)
            sets=', '.join(f'"{x}"=%s' for x in vals)
            cur.execute(f'UPDATE ventas SET {sets} WHERE "numfactur"::text=%s',[*vals.values(),numfactur])
            db.commit()
            return JsonResponse({'ok':True})
    except Exception as exc:
        try: db.rollback()
        except Exception: pass
        return JsonResponse({'ok':False,'error':f'Error actualizando venta: {type(exc).__name__}: {exc}'},status=500)


@admin_required
def admin_compras(request, ruc):
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    db = _cliente_db(cliente)

    try:
        with db.cursor() as cursor:
            cursor.execute(
                f"SELECT COUNT(*) FROM ({_compras_base_sql()}) compras_reporte WHERE {where}",
                params,
            )
            total_registros = cursor.fetchone()[0]
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    try:
        pagina = max(1, int(request.GET.get('pagina', '1')))
    except (TypeError, ValueError):
        pagina = 1

    por_pagina = 50
    total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)
    pagina = min(pagina, total_paginas)
    offset = (pagina - 1) * por_pagina

    try:
        filas = _compras_query(where, params.copy(), cliente, por_pagina, offset, include_numcompra=True)
        resumen = _compras_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    return render(request, 'admin/compras.html', {
        'cliente': cliente,
        'filas': filas,
        'resumen': resumen,
        'filtros': filtros,
        'pagina': pagina,
        'total_paginas': total_paginas,
        'total_registros': total_registros,
    })


@admin_required
def admin_compras_editar(request, ruc):
    """Carga y actualiza una compra desde el módulo administrativo."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    numcompra = request.POST.get('numcompra', '').strip() if request.method == 'POST' else request.GET.get('numcompra', '').strip()
    autorizacion = request.POST.get('autorizacion', '').strip() if request.method == 'POST' else request.GET.get('autorizacion', '').strip()
    proveedor_ruc = request.POST.get('proveedor_ruc', '').strip() if request.method == 'POST' else request.GET.get('proveedor_ruc', '').strip()
    documento = request.POST.get('documento', '').strip() if request.method == 'POST' else request.GET.get('documento', '').strip()

    try:
        # Conexión local al PostgreSQL de la base del cliente.
        # Se construye aquí para que el módulo administrativo no dependa
        # de la función privada de views.py.
        alias = f'cliente_{cliente.ruccedcli}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = cliente.ruccedcli
            connections.databases[alias] = base
        db = connections[alias]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error conectando a la base del cliente: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor_ctx = db.cursor()
        cursor = cursor_ctx.__enter__()
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error abriendo cursor de compras: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor.execute("""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema='public' AND table_name='comprasnue'
            ORDER BY ordinal_position
        """)
        columnas = [r[0] for r in cursor.fetchall()]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error consultando estructura de comprasnue: {type(exc).__name__}: {exc}'
        }, status=500)

    if not columnas:
        return JsonResponse({'ok': False, 'error': 'No se encontró la tabla comprasnue.'}, status=404)
    
    cols = set(columnas)
    # Identifica, sin asumir un nombre único, los campos usados por instalaciones antiguas.
    mes_col = next((x for x in ('mesdeclaracion','mesdeclara','mesdec','mes_declaracion','mes') if x in cols), None)
    anio_col = next((x for x in ('aniodeclaracion','aniodeclara','aniodec','anio_declaracion','anio','ano','año') if x in cols), None)
    
    select_cols = [
        'numcompra','fecemi','ruccedprovee','nomprovee','tipcom','numest','numptoemi','numsec','numaut',
        'baseimpnoobj','baseimpiva0','baseexenta','baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    select_cols = [x for x in select_cols if x in cols]
    if mes_col: select_cols.append(mes_col)
    if anio_col: select_cols.append(anio_col)
    
    if 'numcompra' not in cols:
        return JsonResponse({'ok': False, 'error': 'El campo numcompra no existe en comprasnue.'}, status=500)
    if numcompra:
        where = '"numcompra"::text=%s'
        params = [numcompra]
    elif proveedor_ruc and documento:
        partes = documento.split('-')
        if len(partes) != 3:
            return JsonResponse({'ok': False, 'error': 'Número de factura inválido.'}, status=400)
        numest, numptoemi, numsec = [p.strip() for p in partes]
        where = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
        params = [proveedor_ruc, numest, numptoemi, numsec]
    else:
        return JsonResponse({'ok': False, 'error': 'No se recibió el identificador de la compra.'}, status=400)
    
    try:
        select_sql = f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where} LIMIT 1'
        cursor.execute(select_sql, params)
        row = cursor.fetchone()
    
        if not row and numcompra and proveedor_ruc and documento:
            partes = documento.split('-')
            if len(partes) == 3:
                numest, numptoemi, numsec = [p.strip() for p in partes]
                where_fallback = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
                params_fallback = [proveedor_ruc, numest, numptoemi, numsec]
                cursor.execute(
                    f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where_fallback} LIMIT 1',
                    params_fallback
                )
                row = cursor.fetchone()
    except Exception as exc:
        db.rollback()
        return JsonResponse({'ok': False, 'error': f'Error consultando la compra: {type(exc).__name__}: {exc}'}, status=500)
    if not row:
        return JsonResponse({
            'ok': False,
            'error': f'No se encontró la compra. numcompra recibido: {numcompra or "(vacío)"}; RUC: {proveedor_ruc or "(vacío)"}; factura: {documento or "(vacía)"}'
        }, status=404)
    
    data = dict(zip(select_cols, row))
    if request.method == 'GET':
        for k, v in list(data.items()):
            if hasattr(v, 'isoformat'):
                data[k] = v.isoformat()
            elif v is None:
                data[k] = ''
            else:
                data[k] = str(v)
        data['_mes_col'] = mes_col or ''
        data['_anio_col'] = anio_col or ''
        return JsonResponse({'ok': True, 'compra': data})
    
    editable = [
        'fecemi','tipcom','numaut','baseimpnoobj','baseimpiva0','baseexenta',
        'baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    values = {}
    for field in editable:
        if field in cols and field in request.POST:
            values[field] = request.POST.get(field, '').strip() or None
    # IVA se recalcula en servidor a partir de los subtotales.
    tasas = {'baseimpiva5': ('montoiva5', 0.05), 'baseimpiva8': ('montoiva8', 0.08),
             'baseimpiva12': ('montoiva12', 0.12), 'baseimpiva14': ('montoiva14', 0.14),
             'baseimpiva15': ('montoiva15', 0.15)}
    for base, (iva, tasa) in tasas.items():
        if base in cols:
            try:
                base_val = float(values.get(base) or 0)
                values[iva] = round(base_val * tasa, 2)
            except (TypeError, ValueError):
                values[iva] = 0
    
    if mes_col and 'mes_declaracion' in request.POST:
        values[mes_col] = request.POST.get('mes_declaracion') or None
    if anio_col and 'anio_declaracion' in request.POST:
        values[anio_col] = request.POST.get('anio_declaracion') or None
    
    if not values:
        return JsonResponse({'ok': False, 'error': 'No hay cambios para guardar.'}, status=400)
    
    sets = ', '.join(f'"{k}"=%s' for k in values)
    cursor.execute(f'UPDATE comprasnue SET {sets} WHERE {where}', [*values.values(), *params])
    db.commit()
    return JsonResponse({'ok': True, 'message': 'Compra actualizada correctamente.'})


@admin_required
def admin_compras_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compras'
    ws.append([label for _, label in COMPRAS_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')
    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', '', 'RESUMEN',
               resumen['bases_sin_iva'], resumen['bases_con_iva'], resumen['iva'],
               resumen['total'], resumen['retiva'], '', resumen['retrenta'], ''])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    widths = [7, 38, 17, 10, 13, 25, 18, 16, 16, 13, 16, 13, 11, 16, 18]
    for i, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.xlsx"'
    return response


@admin_required
def admin_compras_pdf(request, ruc):
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .views import _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('AdminComprasPDFTitulo', parent=styles['Title'], fontName='Helvetica-Bold',
                            fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3)
    cab = ParagraphStyle('AdminComprasPDFCab', parent=styles['Normal'], fontName='Helvetica-Bold',
                         fontSize=8, leading=10, alignment=TA_CENTER)
    tabla = ParagraphStyle('AdminComprasPDFTabla', parent=styles['Normal'], fontName='Helvetica',
                           fontSize=5, leading=5.8, alignment=TA_CENTER, wordWrap='CJK')
    izq = ParagraphStyle('AdminComprasPDFIzq', parent=tabla, alignment=0)
    enc = ParagraphStyle('AdminComprasPDFEnc', parent=tabla, fontName='Helvetica-Bold',
                         textColor=colors.white, leading=6)

    elements = [
        Paragraph('REPORTE DE COMPRAS - ADMINISTRATIVO', titulo),
        Paragraph(f"DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}", cab),
        Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}", cab),
        Spacer(1, 8),
    ]

    data = [[Paragraph(label, enc) for _, label in COMPRAS_COLUMNS]]
    for row in filas:
        values = []
        for i, value in enumerate(row):
            if i in (1, 2, 3, 4, 5, 6, 12, 14):
                values.append(Paragraph(str(value or ''), izq if i == 1 else tabla))
            else:
                values.append(Paragraph(f"{float(value or 0):.2f}", tabla))
        data.append(values)

    data.append(['', '', '', '', '', '', '', f"{resumen['bases_sin_iva']:.2f}",
                 f"{resumen['bases_con_iva']:.2f}", f"{resumen['iva']:.2f}",
                 f"{resumen['total']:.2f}", f"{resumen['retiva']:.2f}", '',
                 f"{resumen['retrenta']:.2f}", ''])

    table = Table(data, repeatRows=1,
                  colWidths=[22, 105, 65, 38, 48, 70, 70, 50, 50, 40, 50, 42, 38, 48, 52])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (7, 1), (11, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.pdf"'
    return response


@admin_required
def admin_clientes(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    clientes = Cliente.objects.all().order_by('nomclient')

    if query:
        from django.db.models import Q
        clientes = clientes.filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if estado == 'activos':
        clientes = clientes.filter(activo=True)
    elif estado == 'inactivos':
        clientes = clientes.filter(activo=False)

    if dia != 'todos':
        clientes = clientes.filter(diadeclaracion=dia)

    if tipdec:
        clientes = clientes.filter(tipdec__iexact=tipdec)

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/clientes.html',
        {
            'clientes': clientes,
            'query': query,
            'estado': estado,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
        },
    )


@admin_required
def admin_suscriptores(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()

    suscriptores = Suscriptor.objects.all()

    if query:
        from django.db.models import Q
        suscriptores = suscriptores.filter(
            Q(nombre__icontains=query) |
            Q(email__icontains=query)
        )

    if estado == 'activos':
        suscriptores = suscriptores.filter(activo=True)
    elif estado == 'inactivos':
        suscriptores = suscriptores.filter(activo=False)

    return render(
        request,
        'admin/suscriptores.html',
        {
            'suscriptores': suscriptores,
            'query': query,
            'estado': estado,
            'suscriptores_total': Suscriptor.objects.count(),
            'suscriptores_activos': Suscriptor.objects.filter(activo=True).count(),
        },
    )


@admin_required
def admin_contrasenas(request):
    """Listado administrativo de clientes activos y sus credenciales tributarias/laborales."""
    query = request.GET.get('q', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    # Al entrar se muestran únicamente clientes activos.
    # Cuando se realiza una búsqueda, se consulta toda la base de datos,
    # incluyendo clientes inactivos.
    clientes_qs = Cliente.objects.filter(activo=True)

    if query:
        from django.db.models import Q
        clientes_qs = Cliente.objects.all().filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if dia != 'todos':
        clientes_qs = clientes_qs.filter(diadeclaracion=str(dia))

    if tipdec:
        clientes_qs = clientes_qs.filter(semensual__iexact=tipdec)

    clientes_qs = clientes_qs.order_by('nomclient')

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/contrasenas.html',
        {
            'clientes': clientes_qs,
            'pestanas': [
                {
                    'dia': dia_item,
                    'clientes': clientes_qs.filter(diadeclaracion=str(dia_item)),
                }
                for dia_item in dias
            ],
            'clientes_busqueda': clientes_qs if query else Cliente.objects.none(),
            'query': query,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


@admin_required
def admin_cliente(request, ruc):
    cliente = get_object_or_404(Cliente, pk=ruc)

    if request.method == 'POST' and request.POST.get('accion') == 'guardar_datos_cliente':
        cliente.dirclient = request.POST.get('dirclient', '').strip()
        cliente.teldomcli = request.POST.get('teldomcli', '').strip()
        cliente.teloficli = request.POST.get('teloficli', '').strip()
        cliente.telcelcli = request.POST.get('telcelcli', '').strip()
        cliente.corelectr = request.POST.get('corelectr', '').strip()
        cliente.ocuclient = request.POST.get('ocuclient', '').strip()
        cliente.save(update_fields=[
            'dirclient', 'teldomcli', 'teloficli',
            'telcelcli', 'corelectr', 'ocuclient',
        ])
        messages.success(request, 'Datos básicos del cliente actualizados correctamente.')
        return redirect('admin_cliente', ruc=cliente.ruccedcli)

    usuarios = (
        UsuarioCliente.objects
        .filter(cliente=cliente)
        .select_related('usuario')
        .order_by('usuario__username')
    )

    db_name = str(cliente.ruccedcli).strip()
    db_status = 'No verificada'
    db_error = ''
    conciliacion = {
        'anio': 0,
        'anios': [],
        'meses': [],
        'totales': {
            'ventas': 0,
            'compras': 0,
            'retrenta': 0,
            'retiva': 0,
            'resultado': 0,
            'retenciones': 0,
        },
    }

    # La conexión y la consulta de conciliación se manejan por separado.
    # Un error en una consulta no debe hacer aparecer la base como desconectada.
    try:
        alias = f'cliente_{db_name}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = db_name
            connections.databases[alias] = base
        client_connection = connections[alias]
        client_connection.ensure_connection()
        db_status = 'Conectada'
    except Exception as exc:
        db_status = 'No disponible'
        db_error = str(exc)

    if db_status == 'Conectada':
        try:
            conciliacion = _admin_conciliacion_cliente(cliente, request.GET.get('anio'))
        except Exception as exc:
            db_error = f'Error consultando datos: {exc}'

    return render(
        request,
        'admin/cliente.html',
        {
            'cliente': cliente,
            'usuarios': usuarios,
            'db_name': db_name,
            'db_status': db_status,
            'db_error': db_error,
            'conciliacion': conciliacion,
        },
    )


def _conta_sincronizar_compras(usuario, ruc, anio, mes, tipo_comprobante=1):
    """Llama a Conta desde el servidor; el navegador nunca recibe el secreto."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
        'tipo_comprobante': int(tipo_comprobante),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc




def _conta_sincronizar_ventas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de facturas emitidas hacia ventas."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/ventas/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_sri_status(usuario, job_id):
    """Consulta el estado de un trabajo SRI ya iniciado en Conta."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/estado/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='GET',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


def _conta_sri_cancelar(usuario, job_id):
    """Solicita a Conta detener un trabajo SRI."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/cancelar/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_chat(usuario, mensaje, conversation_id='admin'):
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'mensaje': mensaje, 'conversation_id': conversation_id or 'admin'}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/chat',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=180) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


@admin_required
def admin_conta_sri_cancelar(request):
    """Solicita la cancelación de un trabajo SRI en Conta."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.POST.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_cancelar(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta_sri_status(request):
    """Endpoint AJAX del panel para consultar trabajos SRI en segundo plano."""
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.GET.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_status(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta(request):
    """Panel de Conta dentro del administrador autenticado de TotalCounts."""
    from datetime import datetime

    clientes = Cliente.objects.filter(activo=True).order_by('nomclient')
    ruc = request.POST.get('ruc', '').strip() if request.method == 'POST' else request.GET.get('ruc', '').strip()
    anio = request.POST.get('anio', '').strip() if request.method == 'POST' else request.GET.get('anio', '').strip()
    mes = request.POST.get('mes', '').strip() if request.method == 'POST' else request.GET.get('mes', '').strip()

    if not anio:
        anio = str(datetime.now().year)
    if not mes:
        mes = str(datetime.now().month)

    resultado = None
    sri_job = None
    respuesta_ia = None
    error = ''

    if request.method == 'POST' and request.POST.get('accion') == 'chat':
        mensaje = request.POST.get('mensaje', '').strip()
        if not mensaje:
            error = 'Escriba una consulta para Conta.'
        else:
            try:
                respuesta_ia = _conta_chat(
                    usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
                    mensaje=mensaje,
                    conversation_id=request.session.session_key or 'admin',
                )
            except RuntimeError as exc:
                error = str(exc)

    elif request.method == 'POST' and request.POST.get('accion') in ('sincronizar_compras', 'sincronizar_ventas', 'validar_ventas'):
        cliente = clientes.filter(ruccedcli=ruc).first()
        if cliente is None:
            error = 'Seleccione un cliente activo válido.'
        else:
            try:
                usuario = request.session.get(ADMIN_USERNAME_KEY, 'ADMIN')
                if request.POST.get('accion') == 'sincronizar_ventas':
                    resultado = _conta_sincronizar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de ventas emitidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'validar_ventas':
                    resultado = _conta_validar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La validación de ventas contra el SRI fue iniciada en segundo plano.')
                else:
                    resultado = _conta_sincronizar_compras(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                        tipo_comprobante=int(request.POST.get('tipo_comprobante', '1')),
                    )
                    messages.success(request, 'La sincronización de compras fue iniciada en segundo plano.')
                sri_job = resultado.get('job_id')
            except (ValueError, RuntimeError) as exc:
                error = str(exc)

    return render(
        request,
        'admin/conta.html',
        {
            'clientes': clientes,
            'ruc': ruc,
            'anio': anio,
            'mes': mes,
            'resultado': resultado,
            'sri_job': sri_job,
            'respuesta_ia': respuesta_ia,
            'error': error,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


def _admin_compras_filtros(request):
    hoy = timezone.localdate()
    primer_dia_mes = hoy.replace(day=1)
    fecha_desde = request.GET.get('fecha_desde', '').strip() or primer_dia_mes.strftime('%Y-%m-%d')
    fecha_hasta = request.GET.get('fecha_hasta', '').strip() or hoy.strftime('%Y-%m-%d')
    proveedor = request.GET.get('proveedor', '').strip()

    where = []
    params = []
    try:
        datetime.strptime(fecha_desde, '%Y-%m-%d')
        where.append('fecemi::date >= %s::date')
        params.append(fecha_desde)
    except ValueError:
        fecha_desde = ''

    try:
        datetime.strptime(fecha_hasta, '%Y-%m-%d')
        where.append("fecemi::date < (%s::date + INTERVAL '1 day')")
        params.append(fecha_hasta)
    except ValueError:
        fecha_hasta = ''

    if proveedor:
        where.append('(ruccedprovee ILIKE %s OR nomprovee ILIKE %s)')
        params.extend([f'%{proveedor}%', f'%{proveedor}%'])

    return (' AND '.join(where) if where else '1=1'), params, {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'proveedor': proveedor,
    }


@admin_required
def admin_ventas(request, ruc):
    from .views import _cliente_db, _ventas_base_sql, _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_ventas_filtros(request)
    try:
        base = _ventas_base_sql()
        with _cliente_db(cliente).cursor() as cursor:
            cursor.execute(f"SELECT COUNT(*) FROM ({base}) ventas_reporte WHERE {where}", params)
            total_registros = cursor.fetchone()[0]
        try:
            pagina = max(1, int(request.GET.get('pagina','1')))
        except ValueError:
            pagina = 1
        por_pagina = 50
        filas = _ventas_query(where, params.copy(), cliente, por_pagina, (pagina-1)*por_pagina)
        resumen = _ventas_resumen(where, params.copy(), cliente)
        total_paginas = max(1, (total_registros + por_pagina - 1)//por_pagina)
    except Exception as exc:
        return HttpResponse(f"Error en reporte de ventas: {type(exc).__name__}: {exc}", status=500, content_type='text/plain; charset=utf-8')
    return render(request, 'admin/ventas.html', {
        'cliente': cliente, 'filas': filas, 'resumen': resumen, 'filtros': filtros,
        'pagina': pagina, 'total_paginas': total_paginas, 'total_registros': total_registros,
        'ventas_columns': VENTAS_COLUMNS,
    })


def _admin_ventas_filtros(request):
    hoy = timezone.localdate()
    primer_dia = hoy.replace(day=1)
    desde = request.GET.get('fecha_desde','').strip() or primer_dia.strftime('%Y-%m-%d')
    hasta = request.GET.get('fecha_hasta','').strip() or hoy.strftime('%Y-%m-%d')
    cliente_busqueda = request.GET.get('cliente_busqueda','').strip()
    where=[]; params=[]
    try:
        datetime.strptime(desde,'%Y-%m-%d'); where.append('fecfactur::date >= %s::date'); params.append(desde)
    except ValueError: desde=''
    try:
        datetime.strptime(hasta,'%Y-%m-%d'); where.append("fecfactur::date < (%s::date + INTERVAL '1 day')"); params.append(hasta)
    except ValueError: hasta=''
    if cliente_busqueda:
        where.append('(ruccedcli ILIKE %s OR nomcli ILIKE %s)')
        params.extend([f'%{cliente_busqueda}%',f'%{cliente_busqueda}%'])
    return (' AND '.join(where) if where else '1=1'), params, {'fecha_desde':desde,'fecha_hasta':hasta,'cliente_busqueda':cliente_busqueda}


@admin_required
def admin_ventas_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _cliente_db, _ventas_query, _ventas_resumen, VENTAS_COLUMNS
    cliente=get_object_or_404(Cliente,pk=ruc); where,params,filtros=_admin_ventas_filtros(request)
    filas=_ventas_query(where,params.copy(),cliente); resumen=_ventas_resumen(where,params.copy(),cliente)
    wb=Workbook(); ws=wb.active; ws.title='Ventas'
    ws.append([x[1] for x in VENTAS_COLUMNS])
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='21333E'); cell.alignment=Alignment(horizontal='center')
    for row in filas: ws.append(list(row))
    ws.append([]); ws.append(['','','','','','RESUMEN',resumen['base0'],resumen['baseiva'],resumen['iva'],resumen['total'],resumen['retiva'],resumen['retrenta']])
    buffer=BytesIO(); wb.save(buffer)
    response=HttpResponse(buffer.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); response['Content-Disposition']='attachment; filename="ventas_administrativo.xlsx"'; return response


@admin_required
def admin_ventas_editar(request, ruc):
    from .views import _cliente_db
    cliente=get_object_or_404(Cliente,pk=ruc)
    numfactur=(request.POST.get('numfactur') if request.method=='POST' else request.GET.get('numfactur','')).strip()
    db=_cliente_db(cliente)
    try:
        with db.cursor() as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_schema='public' AND table_name='ventas' ORDER BY ordinal_position")
            cols=[x[0] for x in cur.fetchall()]
            if 'numfactur' not in cols: return JsonResponse({'ok':False,'error':'El campo numfactur no existe en ventas.'},status=500)
            wanted=['fecfactur','numfactur','autorizacion','basenoobj','baseiva0','baseiva12','iva','retiva','retrenta','numret','autret','codret']
            select=[x for x in wanted if x in cols]
            cur.execute(f'SELECT {",".join(chr(34)+x+chr(34) for x in select)} FROM ventas WHERE "numfactur"::text=%s LIMIT 1',[numfactur])
            row=cur.fetchone()
            if not row: return JsonResponse({'ok':False,'error':f'No se encontró la factura {numfactur}.'},status=404)
            data=dict(zip(select,row))
            if request.method=='GET':
                for k,v in data.items(): data[k]='' if v is None else (v.isoformat() if hasattr(v,'isoformat') else str(v))
                return JsonResponse({'ok':True,'venta':data})
            editable=[x for x in ['fecfactur','autorizacion','basenoobj','baseiva0','baseiva12','iva','retiva','retrenta','numret','autret','codret'] if x in cols and x in request.POST]
            vals={x:(request.POST.get(x,'').strip() or None) for x in editable}
            if not vals: return JsonResponse({'ok':False,'error':'No hay cambios para guardar.'},status=400)
            sets=', '.join(f'"{x}"=%s' for x in vals)
            cur.execute(f'UPDATE ventas SET {sets} WHERE "numfactur"::text=%s',[*vals.values(),numfactur])
            db.commit()
            return JsonResponse({'ok':True})
    except Exception as exc:
        try: db.rollback()
        except Exception: pass
        return JsonResponse({'ok':False,'error':f'Error actualizando venta: {type(exc).__name__}: {exc}'},status=500)


@admin_required
def admin_compras(request, ruc):
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    db = _cliente_db(cliente)

    try:
        with db.cursor() as cursor:
            cursor.execute(
                f"SELECT COUNT(*) FROM ({_compras_base_sql()}) compras_reporte WHERE {where}",
                params,
            )
            total_registros = cursor.fetchone()[0]
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    try:
        pagina = max(1, int(request.GET.get('pagina', '1')))
    except (TypeError, ValueError):
        pagina = 1

    por_pagina = 50
    total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)
    pagina = min(pagina, total_paginas)
    offset = (pagina - 1) * por_pagina

    try:
        filas = _compras_query(where, params.copy(), cliente, por_pagina, offset, include_numcompra=True)
        resumen = _compras_resumen(where, params.copy(), cliente)
    except Exception as exc:
        return HttpResponse(f'Error consultando compras: {type(exc).__name__}: {exc}', status=500,
                            content_type='text/plain; charset=utf-8')

    return render(request, 'admin/compras.html', {
        'cliente': cliente,
        'filas': filas,
        'resumen': resumen,
        'filtros': filtros,
        'pagina': pagina,
        'total_paginas': total_paginas,
        'total_registros': total_registros,
    })


@admin_required
def admin_compras_editar(request, ruc):
    """Carga y actualiza una compra desde el módulo administrativo."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    numcompra = request.POST.get('numcompra', '').strip() if request.method == 'POST' else request.GET.get('numcompra', '').strip()
    autorizacion = request.POST.get('autorizacion', '').strip() if request.method == 'POST' else request.GET.get('autorizacion', '').strip()
    proveedor_ruc = request.POST.get('proveedor_ruc', '').strip() if request.method == 'POST' else request.GET.get('proveedor_ruc', '').strip()
    documento = request.POST.get('documento', '').strip() if request.method == 'POST' else request.GET.get('documento', '').strip()

    try:
        # Conexión local al PostgreSQL de la base del cliente.
        # Se construye aquí para que el módulo administrativo no dependa
        # de la función privada de views.py.
        alias = f'cliente_{cliente.ruccedcli}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = cliente.ruccedcli
            connections.databases[alias] = base
        db = connections[alias]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error conectando a la base del cliente: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor_ctx = db.cursor()
        cursor = cursor_ctx.__enter__()
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error abriendo cursor de compras: {type(exc).__name__}: {exc}'
        }, status=500)

    try:
        cursor.execute("""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema='public' AND table_name='comprasnue'
            ORDER BY ordinal_position
        """)
        columnas = [r[0] for r in cursor.fetchall()]
    except Exception as exc:
        return JsonResponse({
            'ok': False,
            'error': f'Error consultando estructura de comprasnue: {type(exc).__name__}: {exc}'
        }, status=500)

    if not columnas:
        return JsonResponse({'ok': False, 'error': 'No se encontró la tabla comprasnue.'}, status=404)
    
    cols = set(columnas)
    # Identifica, sin asumir un nombre único, los campos usados por instalaciones antiguas.
    mes_col = next((x for x in ('mesdeclaracion','mesdeclara','mesdec','mes_declaracion','mes') if x in cols), None)
    anio_col = next((x for x in ('aniodeclaracion','aniodeclara','aniodec','anio_declaracion','anio','ano','año') if x in cols), None)
    
    select_cols = [
        'numcompra','fecemi','ruccedprovee','nomprovee','tipcom','numest','numptoemi','numsec','numaut',
        'baseimpnoobj','baseimpiva0','baseexenta','baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    select_cols = [x for x in select_cols if x in cols]
    if mes_col: select_cols.append(mes_col)
    if anio_col: select_cols.append(anio_col)
    
    if 'numcompra' not in cols:
        return JsonResponse({'ok': False, 'error': 'El campo numcompra no existe en comprasnue.'}, status=500)
    if numcompra:
        where = '"numcompra"::text=%s'
        params = [numcompra]
    elif proveedor_ruc and documento:
        partes = documento.split('-')
        if len(partes) != 3:
            return JsonResponse({'ok': False, 'error': 'Número de factura inválido.'}, status=400)
        numest, numptoemi, numsec = [p.strip() for p in partes]
        where = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
        params = [proveedor_ruc, numest, numptoemi, numsec]
    else:
        return JsonResponse({'ok': False, 'error': 'No se recibió el identificador de la compra.'}, status=400)
    
    try:
        select_sql = f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where} LIMIT 1'
        cursor.execute(select_sql, params)
        row = cursor.fetchone()
    
        if not row and numcompra and proveedor_ruc and documento:
            partes = documento.split('-')
            if len(partes) == 3:
                numest, numptoemi, numsec = [p.strip() for p in partes]
                where_fallback = '"ruccedprovee"::text=%s AND "numest"::text=%s AND "numptoemi"::text=%s AND "numsec"::text=%s'
                params_fallback = [proveedor_ruc, numest, numptoemi, numsec]
                cursor.execute(
                    f'SELECT {", ".join(chr(34)+x+chr(34) for x in select_cols)} FROM comprasnue WHERE {where_fallback} LIMIT 1',
                    params_fallback
                )
                row = cursor.fetchone()
    except Exception as exc:
        db.rollback()
        return JsonResponse({'ok': False, 'error': f'Error consultando la compra: {type(exc).__name__}: {exc}'}, status=500)
    if not row:
        return JsonResponse({
            'ok': False,
            'error': f'No se encontró la compra. numcompra recibido: {numcompra or "(vacío)"}; RUC: {proveedor_ruc or "(vacío)"}; factura: {documento or "(vacía)"}'
        }, status=404)
    
    data = dict(zip(select_cols, row))
    if request.method == 'GET':
        for k, v in list(data.items()):
            if hasattr(v, 'isoformat'):
                data[k] = v.isoformat()
            elif v is None:
                data[k] = ''
            else:
                data[k] = str(v)
        data['_mes_col'] = mes_col or ''
        data['_anio_col'] = anio_col or ''
        return JsonResponse({'ok': True, 'compra': data})
    
    editable = [
        'fecemi','tipcom','numaut','baseimpnoobj','baseimpiva0','baseexenta',
        'baseimpiva5','baseimpiva8','baseimpiva12','baseimpiva14','baseimpiva15',
        'montoiva5','montoiva8','montoiva12','montoiva14','montoiva15',
        'retencioniva10','retencioniva20','retencioniva30','retencioniva70','retencioniva100',
        'codret','valret','numestret','numptoemiret','numsecret'
    ]
    values = {}
    for field in editable:
        if field in cols and field in request.POST:
            values[field] = request.POST.get(field, '').strip() or None
    # IVA se recalcula en servidor a partir de los subtotales.
    tasas = {'baseimpiva5': ('montoiva5', 0.05), 'baseimpiva8': ('montoiva8', 0.08),
             'baseimpiva12': ('montoiva12', 0.12), 'baseimpiva14': ('montoiva14', 0.14),
             'baseimpiva15': ('montoiva15', 0.15)}
    for base, (iva, tasa) in tasas.items():
        if base in cols:
            try:
                base_val = float(values.get(base) or 0)
                values[iva] = round(base_val * tasa, 2)
            except (TypeError, ValueError):
                values[iva] = 0
    
    if mes_col and 'mes_declaracion' in request.POST:
        values[mes_col] = request.POST.get('mes_declaracion') or None
    if anio_col and 'anio_declaracion' in request.POST:
        values[anio_col] = request.POST.get('anio_declaracion') or None
    
    if not values:
        return JsonResponse({'ok': False, 'error': 'No hay cambios para guardar.'}, status=400)
    
    sets = ', '.join(f'"{k}"=%s' for k in values)
    cursor.execute(f'UPDATE comprasnue SET {sets} WHERE {where}', [*values.values(), *params])
    db.commit()
    return JsonResponse({'ok': True, 'message': 'Compra actualizada correctamente.'})


@admin_required
def admin_compras_excel(request, ruc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .views import _cliente_db, _compras_base_sql, _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compras'
    ws.append([label for _, label in COMPRAS_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')
    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', '', 'RESUMEN',
               resumen['bases_sin_iva'], resumen['bases_con_iva'], resumen['iva'],
               resumen['total'], resumen['retiva'], '', resumen['retrenta'], ''])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    widths = [7, 38, 17, 10, 13, 25, 18, 16, 16, 13, 16, 13, 11, 16, 18]
    for i, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.xlsx"'
    return response


@admin_required
def admin_compras_pdf(request, ruc):
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from .views import _compras_query, _compras_resumen, COMPRAS_COLUMNS

    cliente = get_object_or_404(Cliente, pk=ruc)
    where, params, filtros = _admin_compras_filtros(request)
    filas = _compras_query(where, params.copy(), cliente)
    resumen = _compras_resumen(where, params.copy(), cliente)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('AdminComprasPDFTitulo', parent=styles['Title'], fontName='Helvetica-Bold',
                            fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3)
    cab = ParagraphStyle('AdminComprasPDFCab', parent=styles['Normal'], fontName='Helvetica-Bold',
                         fontSize=8, leading=10, alignment=TA_CENTER)
    tabla = ParagraphStyle('AdminComprasPDFTabla', parent=styles['Normal'], fontName='Helvetica',
                           fontSize=5, leading=5.8, alignment=TA_CENTER, wordWrap='CJK')
    izq = ParagraphStyle('AdminComprasPDFIzq', parent=tabla, alignment=0)
    enc = ParagraphStyle('AdminComprasPDFEnc', parent=tabla, fontName='Helvetica-Bold',
                         textColor=colors.white, leading=6)

    elements = [
        Paragraph('REPORTE DE COMPRAS - ADMINISTRATIVO', titulo),
        Paragraph(f"DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}", cab),
        Paragraph(f"{cliente.nomclient} | RUC. {cliente.ruccedcli}", cab),
        Spacer(1, 8),
    ]

    data = [[Paragraph(label, enc) for _, label in COMPRAS_COLUMNS]]
    for row in filas:
        values = []
        for i, value in enumerate(row):
            if i in (1, 2, 3, 4, 5, 6, 12, 14):
                values.append(Paragraph(str(value or ''), izq if i == 1 else tabla))
            else:
                values.append(Paragraph(f"{float(value or 0):.2f}", tabla))
        data.append(values)

    data.append(['', '', '', '', '', '', '', f"{resumen['bases_sin_iva']:.2f}",
                 f"{resumen['bases_con_iva']:.2f}", f"{resumen['iva']:.2f}",
                 f"{resumen['total']:.2f}", f"{resumen['retiva']:.2f}", '',
                 f"{resumen['retrenta']:.2f}", ''])

    table = Table(data, repeatRows=1,
                  colWidths=[22, 105, 65, 38, 48, 70, 70, 50, 50, 40, 50, 42, 38, 48, 52])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (7, 1), (11, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="compras_administrativo.pdf"'
    return response


@admin_required
def admin_clientes(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    clientes = Cliente.objects.all().order_by('nomclient')

    if query:
        from django.db.models import Q
        clientes = clientes.filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if estado == 'activos':
        clientes = clientes.filter(activo=True)
    elif estado == 'inactivos':
        clientes = clientes.filter(activo=False)

    if dia != 'todos':
        clientes = clientes.filter(diadeclaracion=dia)

    if tipdec:
        clientes = clientes.filter(tipdec__iexact=tipdec)

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/clientes.html',
        {
            'clientes': clientes,
            'query': query,
            'estado': estado,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
        },
    )


@admin_required
def admin_suscriptores(request):
    query = request.GET.get('q', '').strip()
    estado = request.GET.get('estado', '').strip()

    suscriptores = Suscriptor.objects.all()

    if query:
        from django.db.models import Q
        suscriptores = suscriptores.filter(
            Q(nombre__icontains=query) |
            Q(email__icontains=query)
        )

    if estado == 'activos':
        suscriptores = suscriptores.filter(activo=True)
    elif estado == 'inactivos':
        suscriptores = suscriptores.filter(activo=False)

    return render(
        request,
        'admin/suscriptores.html',
        {
            'suscriptores': suscriptores,
            'query': query,
            'estado': estado,
            'suscriptores_total': Suscriptor.objects.count(),
            'suscriptores_activos': Suscriptor.objects.filter(activo=True).count(),
        },
    )


@admin_required
def admin_contrasenas(request):
    """Listado administrativo de clientes activos y sus credenciales tributarias/laborales."""
    query = request.GET.get('q', '').strip()
    dia = request.GET.get('dia', 'todos').strip() or 'todos'
    tipdec = request.GET.get('tipdec', '').strip()

    # Al entrar se muestran únicamente clientes activos.
    # Cuando se realiza una búsqueda, se consulta toda la base de datos,
    # incluyendo clientes inactivos.
    clientes_qs = Cliente.objects.filter(activo=True)

    if query:
        from django.db.models import Q
        clientes_qs = Cliente.objects.all().filter(
            Q(nomclient__icontains=query) |
            Q(ruccedcli__icontains=query)
        )

    if dia != 'todos':
        clientes_qs = clientes_qs.filter(diadeclaracion=str(dia))

    if tipdec:
        clientes_qs = clientes_qs.filter(semensual__iexact=tipdec)

    clientes_qs = clientes_qs.order_by('nomclient')

    dias = (10, 12, 14, 16, 18, 20, 22, 24, 26, 28)
    tipos_dec = ('MENSUAL', 'SEMESTRAL', 'ANUAL')

    return render(
        request,
        'admin/contrasenas.html',
        {
            'clientes': clientes_qs,
            'pestanas': [
                {
                    'dia': dia_item,
                    'clientes': clientes_qs.filter(diadeclaracion=str(dia_item)),
                }
                for dia_item in dias
            ],
            'clientes_busqueda': clientes_qs if query else Cliente.objects.none(),
            'query': query,
            'dia': dia,
            'tipdec': tipdec,
            'dias': dias,
            'tipos_dec': tipos_dec,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )


@admin_required
def admin_documentacion(request, ruc):
    """Muestra las categorías de documentación disponibles para el cliente."""
    cliente = get_object_or_404(Cliente, pk=ruc)
    return render(
        request,
        'admin/documentacion.html',
        {
            'cliente': cliente,
            'db_name': str(cliente.ruccedcli).strip(),
        },
    )


@admin_required
def admin_cliente(request, ruc):
    cliente = get_object_or_404(Cliente, pk=ruc)

    if request.method == 'POST' and request.POST.get('accion') == 'guardar_datos_cliente':
        cliente.dirclient = request.POST.get('dirclient', '').strip()
        cliente.teldomcli = request.POST.get('teldomcli', '').strip()
        cliente.teloficli = request.POST.get('teloficli', '').strip()
        cliente.telcelcli = request.POST.get('telcelcli', '').strip()
        cliente.corelectr = request.POST.get('corelectr', '').strip()
        cliente.ocuclient = request.POST.get('ocuclient', '').strip()
        cliente.save(update_fields=[
            'dirclient', 'teldomcli', 'teloficli',
            'telcelcli', 'corelectr', 'ocuclient',
        ])
        messages.success(request, 'Datos básicos del cliente actualizados correctamente.')
        return redirect('admin_cliente', ruc=cliente.ruccedcli)

    usuarios = (
        UsuarioCliente.objects
        .filter(cliente=cliente)
        .select_related('usuario')
        .order_by('usuario__username')
    )

    db_name = str(cliente.ruccedcli).strip()
    db_status = 'No verificada'
    db_error = ''
    conciliacion = {
        'anio': 0,
        'anios': [],
        'meses': [],
        'totales': {
            'ventas': 0,
            'compras': 0,
            'retrenta': 0,
            'retiva': 0,
            'resultado': 0,
            'retenciones': 0,
        },
    }

    # La conexión y la consulta de conciliación se manejan por separado.
    # Un error en una consulta no debe hacer aparecer la base como desconectada.
    try:
        alias = f'cliente_{db_name}'
        if alias not in connections.databases:
            base = settings.DATABASES['default'].copy()
            base['NAME'] = db_name
            connections.databases[alias] = base
        client_connection = connections[alias]
        client_connection.ensure_connection()
        db_status = 'Conectada'
    except Exception as exc:
        db_status = 'No disponible'
        db_error = str(exc)

    if db_status == 'Conectada':
        try:
            conciliacion = _admin_conciliacion_cliente(cliente, request.GET.get('anio'))
        except Exception as exc:
            db_error = f'Error consultando datos: {exc}'

    return render(
        request,
        'admin/cliente.html',
        {
            'cliente': cliente,
            'usuarios': usuarios,
            'db_name': db_name,
            'db_status': db_status,
            'db_error': db_error,
            'conciliacion': conciliacion,
        },
    )


def _conta_sincronizar_compras(usuario, ruc, anio, mes, tipo_comprobante=1):
    """Llama a Conta desde el servidor; el navegador nunca recibe el secreto."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
        'tipo_comprobante': int(tipo_comprobante),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc




def _conta_sincronizar_ventas(usuario, ruc, anio, mes):
    """Inicia en Conta la sincronización de facturas emitidas hacia ventas."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({
        'ruc': ruc, 'anio': int(anio), 'mes': int(mes),
    }).encode('utf-8')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/ventas/sincronizar',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=20) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_sri_status(usuario, job_id):
    """Consulta el estado de un trabajo SRI ya iniciado en Conta."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/compras/estado/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='GET',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


def _conta_sri_cancelar(usuario, job_id):
    """Solicita a Conta detener un trabajo SRI."""
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    req = request.Request(
        f'{base_url}/api/v1/admin/sri/cancelar/{job_id}',
        headers={
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=10) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}')


def _conta_chat(usuario, mensaje, conversation_id='admin'):
    import json
    import os
    from urllib import error, request

    base_url = os.getenv('CONTA_URL', 'http://127.0.0.1:2408').rstrip('/')
    token = os.getenv('CONTA_INTERNAL_TOKEN', '').strip()
    if not token:
        raise RuntimeError('CONTA_INTERNAL_TOKEN no está configurado en TotalCounts.')

    payload = json.dumps({'mensaje': mensaje, 'conversation_id': conversation_id or 'admin'}).encode('utf-8')
    req = request.Request(
        f'{base_url}/api/v1/admin/chat',
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'Accept': 'application/json',
            'X-TotalCounts-Internal': token,
            'X-TotalCounts-User': usuario,
        },
        method='POST',
    )
    try:
        with request.urlopen(req, timeout=180) as response:
            return json.loads(response.read().decode('utf-8'))
    except error.HTTPError as exc:
        body = exc.read().decode('utf-8', errors='replace')
        try:
            detail = json.loads(body).get('detail', body)
        except json.JSONDecodeError:
            detail = body
        raise RuntimeError(f'Conta respondió HTTP {exc.code}: {detail}') from exc
    except error.URLError as exc:
        raise RuntimeError(f'No se pudo conectar con Conta: {exc.reason}') from exc


@admin_required
def admin_conta_sri_cancelar(request):
    """Solicita la cancelación de un trabajo SRI en Conta."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.POST.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_cancelar(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta_sri_status(request):
    """Endpoint AJAX del panel para consultar trabajos SRI en segundo plano."""
    if request.method != 'GET':
        return JsonResponse({'error': 'Método no permitido.'}, status=405)
    job_id = request.GET.get('job_id', '').strip()
    if not job_id:
        return JsonResponse({'error': 'Falta job_id.'}, status=400)
    try:
        resultado = _conta_sri_status(
            usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
            job_id=job_id,
        )
        return JsonResponse(resultado)
    except RuntimeError as exc:
        return JsonResponse({'error': str(exc)}, status=502)


@admin_required
def admin_conta(request):
    """Panel de Conta dentro del administrador autenticado de TotalCounts."""
    from datetime import datetime

    clientes = Cliente.objects.filter(activo=True).order_by('nomclient')
    ruc = request.POST.get('ruc', '').strip() if request.method == 'POST' else request.GET.get('ruc', '').strip()
    anio = request.POST.get('anio', '').strip() if request.method == 'POST' else request.GET.get('anio', '').strip()
    mes = request.POST.get('mes', '').strip() if request.method == 'POST' else request.GET.get('mes', '').strip()

    if not anio:
        anio = str(datetime.now().year)
    if not mes:
        mes = str(datetime.now().month)

    resultado = None
    sri_job = None
    respuesta_ia = None
    error = ''

    if request.method == 'POST' and request.POST.get('accion') == 'chat':
        mensaje = request.POST.get('mensaje', '').strip()
        if not mensaje:
            error = 'Escriba una consulta para Conta.'
        else:
            try:
                respuesta_ia = _conta_chat(
                    usuario=request.session.get(ADMIN_USERNAME_KEY, 'ADMIN'),
                    mensaje=mensaje,
                    conversation_id=request.session.session_key or 'admin',
                )
            except RuntimeError as exc:
                error = str(exc)

    elif request.method == 'POST' and request.POST.get('accion') in ('sincronizar_compras', 'sincronizar_ventas', 'sincronizar_notas_credito_emitidas', 'sincronizar_notas_credito_recibidas', 'sincronizar_retenciones_recibidas', 'validar_ventas'):
        cliente = clientes.filter(ruccedcli=ruc).first()
        if cliente is None:
            error = 'Seleccione un cliente activo válido.'
        else:
            try:
                usuario = request.session.get(ADMIN_USERNAME_KEY, 'ADMIN')
                if request.POST.get('accion') == 'sincronizar_ventas':
                    resultado = _conta_sincronizar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de ventas emitidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'validar_ventas':
                    resultado = _conta_validar_ventas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La validación de ventas contra el SRI fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_notas_credito_recibidas':
                    resultado = _conta_sincronizar_notas_credito_recibidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de notas de crédito recibidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_retenciones_recibidas':
                    resultado = _conta_sincronizar_retenciones_recibidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de retenciones recibidas fue iniciada en segundo plano.')
                elif request.POST.get('accion') == 'sincronizar_notas_credito_emitidas':
                    resultado = _conta_sincronizar_notas_credito_emitidas(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                    )
                    messages.success(request, 'La sincronización de notas de crédito emitidas fue iniciada en segundo plano.')
                else:
                    resultado = _conta_sincronizar_compras(
                        usuario=usuario,
                        ruc=str(cliente.ruccedcli).strip(),
                        anio=int(anio),
                        mes=int(mes),
                        tipo_comprobante=int(request.POST.get('tipo_comprobante', '1')),
                    )
                    messages.success(request, 'La sincronización de compras fue iniciada en segundo plano.')
                sri_job = resultado.get('job_id')
            except (ValueError, RuntimeError) as exc:
                error = str(exc)

    return render(
        request,
        'admin/conta.html',
        {
            'clientes': clientes,
            'ruc': ruc,
            'anio': anio,
            'mes': mes,
            'resultado': resultado,
            'sri_job': sri_job,
            'respuesta_ia': respuesta_ia,
            'error': error,
            'admin_nombre': request.session.get(ADMIN_NAME_KEY, ''),
            'admin_usuario': request.session.get(ADMIN_USERNAME_KEY, ''),
        },
    )
