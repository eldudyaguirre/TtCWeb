from datetime import datetime
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render

from .views import _cliente_db, _cliente_portal


INGRESOS_NC_COLUMNS = [
    ('numero', 'N°'),
    ('cliente', 'CLIENTE'),
    ('ruc', 'RUC'),
    ('fecha', 'FECHA'),
    ('num_nc', 'NUM NC'),
    ('num_aut', 'NUM AUT'),
    ('bases_sin_iva', 'BASES SIN IVA'),
    ('bases_con_iva', 'BASES CON IVA'),
    ('iva', 'IVA'),
    ('total', 'TOTAL'),
    ('fac_mod', 'FAC MOD'),
    ('fecha_factura', 'FECHA FACTURA'),
]


def _ingresos_nc_where(request):
    cliente = _cliente_portal(request)
    if cliente is None:
        return None, [], None

    where = []
    params = []
    hoy = datetime.now().date()
    primer_dia_mes = hoy.replace(day=1)

    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    cliente_busqueda = request.GET.get('cliente_busqueda', '').strip()

    if not fecha_desde:
        fecha_desde = primer_dia_mes.strftime('%Y-%m-%d')
    if not fecha_hasta:
        fecha_hasta = hoy.strftime('%Y-%m-%d')

    if fecha_desde:
        try:
            datetime.strptime(fecha_desde, '%Y-%m-%d')
            where.append("fecncc::date >= %s::date")
            params.append(fecha_desde)
        except ValueError:
            fecha_desde = ''

    if fecha_hasta:
        try:
            datetime.strptime(fecha_hasta, '%Y-%m-%d')
            where.append("fecncc::date < (%s::date + INTERVAL '1 day')")
            params.append(fecha_hasta)
        except ValueError:
            fecha_hasta = ''

    if cliente_busqueda:
        where.append("(ruccedcli ILIKE %s OR nomcli ILIKE %s)")
        params.extend([f'%{cliente_busqueda}%', f'%{cliente_busqueda}%'])

    return " AND ".join(where) if where else "1=1", params, {
        'cliente': cliente,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'cliente_busqueda': cliente_busqueda,
    }


def _ingresos_nc_base_sql():
    return """
        SELECT
            n.fecncc,
            n.nomcli,
            n.ruccedcli,
            n.numnc,
            n.autorizacion,
            n.basenoobj,
            n.baseiva0,
            n.baseiva12,
            n.iva,
            n.numfac,
            n.fecfac
        FROM ncventas n
    """


def _ingresos_nc_query(where, params, cliente, limit=None, offset=None):
    sql = f"""
        SELECT
            ROW_NUMBER() OVER (ORDER BY fecncc::date ASC, num_nc ASC) AS numero,
            nomcli AS cliente,
            ruccedcli AS ruc,
            fecncc AS fecha,
            num_nc,
            num_aut,
            (
                COALESCE(NULLIF(basenoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseiva0::text, ''), '0')::numeric
            ) AS bases_sin_iva,
            COALESCE(NULLIF(baseiva12::text, ''), '0')::numeric AS bases_con_iva,
            COALESCE(NULLIF(iva::text, ''), '0')::numeric AS iva,
            (
                COALESCE(NULLIF(basenoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseiva0::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseiva12::text, ''), '0')::numeric
                + COALESCE(NULLIF(iva::text, ''), '0')::numeric
            ) AS total,
            numfac AS fac_mod,
            fecfac AS fecha_factura
        FROM ({_ingresos_nc_base_sql()}) nc_reporte
        WHERE {where}
        ORDER BY fecncc::date ASC, num_nc ASC
    """
    if limit is not None:
        sql += " LIMIT %s OFFSET %s"
        params = [*params, limit, offset or 0]

    with _cliente_db(cliente).cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.fetchall()


def _ingresos_nc_resumen(where, params, cliente):
    sql = f"""
        SELECT
            COALESCE(SUM(COALESCE(NULLIF(basenoobj::text, ''), '0')::numeric
                + COALESCE(NULLIF(baseiva0::text, ''), '0')::numeric), 0),
            COALESCE(SUM(COALESCE(NULLIF(baseiva12::text, ''), '0')::numeric), 0),
            COALESCE(SUM(COALESCE(NULLIF(iva::text, ''), '0')::numeric), 0)
        FROM ({_ingresos_nc_base_sql()}) nc_reporte
        WHERE {where}
    """
    with _cliente_db(cliente).cursor() as cursor:
        cursor.execute(sql, params)
        row = cursor.fetchone()

    bases_sin, bases_con, iva = [float(value or 0) for value in row]
    return {
        'bases_sin_iva': bases_sin,
        'bases_con_iva': bases_con,
        'iva': iva,
        'total': bases_sin + bases_con + iva,
    }


def _ingresos_nc_datos(request):
    where, params, filtros = _ingresos_nc_where(request)
    if where is None:
        return None, [], None, None
    cliente = filtros['cliente']
    return filtros, _ingresos_nc_query(where, params, cliente), _ingresos_nc_resumen(where, params, cliente), where


@login_required
def ingresos_notas_credito(request):
    try:
        filtros, _, resumen, _ = _ingresos_nc_datos(request)
        if filtros is None:
            return redirect('portal')

        try:
            pagina = max(1, int(request.GET.get('pagina', '1')))
        except ValueError:
            pagina = 1

        por_pagina = 50
        where, params, _ = _ingresos_nc_where(request)
        count_sql = f"SELECT COUNT(*) FROM ({_ingresos_nc_base_sql()}) nc_reporte WHERE {where}"

        with _cliente_db(filtros['cliente']).cursor() as cursor:
            cursor.execute(count_sql, params)
            total_registros = cursor.fetchone()[0]

        offset = (pagina - 1) * por_pagina
        filas = _ingresos_nc_query(where, params, filtros['cliente'], por_pagina, offset)
        total_paginas = max(1, (total_registros + por_pagina - 1) // por_pagina)

        return render(request, 'ingresos-notas-credito.html', {
            'cliente': filtros['cliente'],
            'filas': filas,
            'resumen': resumen,
            'filtros': filtros,
            'pagina': pagina,
            'total_paginas': total_paginas,
            'total_registros': total_registros,
        })
    except Exception as exc:
        return HttpResponse(
            f"Error en reporte de notas de crédito de ingresos: {type(exc).__name__}: {exc}",
            status=500,
            content_type='text/plain; charset=utf-8',
        )


@login_required
def ingresos_notas_credito_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    filtros, filas, resumen, _ = _ingresos_nc_datos(request)
    if filtros is None:
        return redirect('portal')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle('IngresosNCTitulo', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=14, leading=16, alignment=TA_CENTER, spaceAfter=3)
    cabecera = ParagraphStyle('IngresosNCCabecera', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, alignment=TA_CENTER, spaceAfter=2)
    tabla = ParagraphStyle('IngresosNCTabla', parent=styles['Normal'], fontName='Helvetica', fontSize=5.5, leading=6.3, alignment=TA_CENTER, wordWrap='CJK')
    encabezado = ParagraphStyle('IngresosNCEncabezado', parent=tabla, fontName='Helvetica-Bold', textColor=colors.white, leading=6.5)

    elements = [
        Paragraph('REPORTE DE NOTAS DE CRÉDITO', titulo),
        Paragraph(f"NOTAS DE CRÉDITO DESDE {filtros['fecha_desde'] or '—'} A {filtros['fecha_hasta'] or '—'}", cabecera),
        Paragraph(f"{filtros['cliente'].nomclient} | RUC. {filtros['cliente'].ruccedcli}", cabecera),
        Spacer(1, 10),
    ]

    data = [[Paragraph(label, encabezado) for _, label in INGRESOS_NC_COLUMNS]]
    for row in filas:
        formatted = []
        for index, value in enumerate(row):
            if index in (1, 2, 3, 4, 5, 10, 11):
                formatted.append(Paragraph(str(value or ''), tabla))
            else:
                formatted.append(f"{float(value or 0):.2f}")
        data.append(formatted)

    data.append(['', '', '', '', '', '', f"{resumen['bases_sin_iva']:.2f}", f"{resumen['bases_con_iva']:.2f}", f"{resumen['iva']:.2f}", f"{resumen['total']:.2f}", '', ''])

    table = Table(data, repeatRows=1, colWidths=[22, 100, 68, 48, 70, 78, 58, 58, 45, 55, 65, 60])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#21333e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 5.5),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#d8e0e3')),
        ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
        ('ALIGN', (6, 1), (9, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#eef5f5')),
    ]))
    elements.append(table)
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_notas_credito_ingresos.pdf"'
    return response


@login_required
def ingresos_notas_credito_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    filtros, filas, resumen, _ = _ingresos_nc_datos(request)
    if filtros is None:
        return redirect('portal')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Notas de crédito'
    ws.append([label for _, label in INGRESOS_NC_COLUMNS])

    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='21333E')
        cell.alignment = Alignment(horizontal='center')

    for row in filas:
        ws.append(list(row))

    ws.append([])
    ws.append(['', '', '', '', '', '', 'RESUMEN'])
    ws.cell(ws.max_row, 7, resumen['bases_sin_iva'])
    ws.cell(ws.max_row, 8, resumen['bases_con_iva'])
    ws.cell(ws.max_row, 9, resumen['iva'])
    ws.cell(ws.max_row, 10, resumen['total'])

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [7, 35, 17, 13, 22, 28, 16, 16, 13, 16, 22, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte_notas_credito_ingresos.xlsx"'
    return response
